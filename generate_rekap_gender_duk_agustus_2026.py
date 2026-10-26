#!/usr/bin/env python3
"""Create an auditable gender recap from the 1 August 2026 Bapenda DUK."""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SOURCE_CANDIDATES = [
    Path("/Users/alif/Downloads/DUK BAPENDA KOTA MEDAN 2026 1 Agustus 2026 (3).xlsx"),
    Path("/Users/alif/Downloads/DUK BAPENDA KOTA MEDAN 2026 1 Agustus 2026 (2).xlsx"),
    Path("/Users/alif/Downloads/DUK BAPENDA KOTA MEDAN 2026 1 Agustus 2026.xlsx"),
]
SOURCE = next((candidate for candidate in SOURCE_CANDIDATES if candidate.exists()), SOURCE_CANDIDATES[0])
OUTPUT = Path(__file__).resolve().parent / "REKAP JUMLAH PEGAWAI PRIA DAN WANITA - DUK 1 AGUSTUS 2026.xlsx"

MAIN_SHEET = "DUK ASN BAPENDA KOTA MEDAN"
FULL_TIME_SHEET = "P3K PENUH WAKTU SESUAI UNIT"
PART_TIME_SHEET = "P3K PARUH WAKTU SESUAI UNIT"
OUTSOURCE_SHEET = "OUTSOURCING"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F3F8FC"
PALE_GREEN = "E2F0D9"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GREY = "666666"
THIN = Side(style="thin", color="B7C9D6")


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def digits(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def normalize_gender(value: object) -> str:
    text = clean(value).upper()
    if text in {"LAKI-LAKI", "LAKI LAKI", "PRIA"}:
        return "Pria"
    if text in {"PEREMPUAN", "WANITA"}:
        return "Wanita"
    return "Belum Terisi"


def nip_set(worksheet) -> set[str]:
    result: set[str] = set()
    for row in worksheet.iter_rows(min_row=7, values_only=True):
        nip = digits(row[2])
        if len(nip) >= 15:
            result.add(nip)
    return result


def load_data() -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    full_time_nips = nip_set(workbook[FULL_TIME_SHEET])
    part_time_nips = nip_set(workbook[PART_TIME_SHEET])

    if full_time_nips & part_time_nips:
        raise ValueError("Terdapat NIP PPPK yang muncul pada daftar penuh waktu dan paruh waktu.")
    if len(full_time_nips) != 46 or len(part_time_nips) != 540:
        raise ValueError("Jumlah PPPK pada sheet unit kerja tidak sesuai kontrol 46 + 540.")

    asn: list[dict[str, object]] = []
    main = workbook[MAIN_SHEET]
    for source_row, row in enumerate(main.iter_rows(min_row=7, values_only=True), start=7):
        nip = digits(row[2])
        status_source = clean(row[29])
        if len(nip) < 15 or status_source not in {"PNS", "CPNS", "PPPK"}:
            continue
        if status_source == "PPPK":
            if nip in full_time_nips:
                status_detail = "PPPK Penuh Waktu"
            elif nip in part_time_nips:
                status_detail = "PPPK Paruh Waktu"
            else:
                raise ValueError(f"NIP PPPK {nip} tidak ditemukan pada sheet rincian PPPK.")
        else:
            status_detail = status_source
        asn.append(
            {
                "name": clean(row[1]),
                "id": nip,
                "status": status_detail,
                "unit": clean(row[16]),
                "gender": normalize_gender(row[32]),
                "source_sheet": MAIN_SHEET,
                "source_row": source_row,
            }
        )

    outsourcing: list[dict[str, object]] = []
    outsource_sheet = workbook[OUTSOURCE_SHEET]
    for source_row, row in enumerate(outsource_sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = clean(row[1])
        nik = digits(row[2])
        if not name or not nik:
            continue
        outsourcing.append(
            {
                "name": name,
                "id": nik,
                "status": "Outsourcing",
                "unit": clean(row[11]),
                "gender": normalize_gender(row[5]),
                "source_sheet": OUTSOURCE_SHEET,
                "source_row": source_row,
            }
        )

    workbook.close()
    if len(asn) != 843 or len({str(person["id"]) for person in asn}) != 843:
        raise ValueError(f"Kontrol ASN gagal: terbaca {len(asn)} data atau terdapat NIP duplikat.")
    if len(outsourcing) != 38 or len({str(person["id"]) for person in outsourcing}) != 38:
        raise ValueError(f"Kontrol outsourcing gagal: terbaca {len(outsourcing)} data atau terdapat NIK duplikat.")
    return asn, outsourcing


def style_cell(cell, *, fill: str | None = None, bold: bool = False, color: str | None = None) -> None:
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=bold, color=color)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def add_summary(workbook: Workbook, asn: list[dict[str, object]], outsourcing: list[dict[str, object]]) -> None:
    worksheet = workbook.active
    worksheet.title = "RINGKASAN"
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:F1")
    worksheet["A1"] = "REKAP JUMLAH PEGAWAI PRIA DAN WANITA"
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet["A1"].font = Font(color=WHITE, bold=True, size=16)
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 34

    worksheet["A3"] = "File sumber"
    worksheet["B3"] = SOURCE.name
    worksheet.merge_cells("B3:F3")
    worksheet["A4"] = "Tanggal basis data"
    worksheet["B4"] = "1 Agustus 2026"
    for cell in (worksheet["A3"], worksheet["A4"]):
        cell.font = Font(bold=True, color=NAVY)

    asn_gender = Counter(str(person["gender"]) for person in asn)
    out_gender = Counter(str(person["gender"]) for person in outsourcing)

    highlights = [
        ("Total ASN", len(asn), PALE_GREEN),
        ("Pria ASN", asn_gender["Pria"], LIGHT_BLUE),
        ("Wanita ASN", asn_gender["Wanita"], YELLOW),
    ]
    for index, (label, value, fill) in enumerate(highlights):
        column = 1 + index * 2
        label_cell = worksheet.cell(6, column, label)
        value_cell = worksheet.cell(6, column + 1, value)
        style_cell(label_cell, fill=fill, bold=True, color=NAVY)
        style_cell(value_cell, fill=fill, bold=True, color=NAVY)
        value_cell.font = Font(bold=True, color=NAVY, size=15)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[6].height = 36

    headers = ["Status Pegawai", "Pria", "Wanita", "Total", "% Pria", "% Wanita"]
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(9, column, header)
        style_cell(cell, fill=BLUE, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    statuses = ["PNS", "CPNS", "PPPK Penuh Waktu", "PPPK Paruh Waktu"]
    status_counts = Counter((str(person["status"]), str(person["gender"])) for person in asn)
    for offset, status in enumerate(statuses, start=1):
        row = 9 + offset
        male = status_counts[(status, "Pria")]
        female = status_counts[(status, "Wanita")]
        total = male + female
        values = [status, male, female, total, male / total if total else 0, female / total if total else 0]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row, column, value)
            style_cell(cell, fill=PALE_BLUE if row % 2 == 0 else None)
        worksheet.cell(row, 5).number_format = "0.00%"
        worksheet.cell(row, 6).number_format = "0.00%"

    total_row = 14
    total_values = ["TOTAL ASN", asn_gender["Pria"], asn_gender["Wanita"], len(asn), asn_gender["Pria"] / len(asn), asn_gender["Wanita"] / len(asn)]
    for column, value in enumerate(total_values, start=1):
        cell = worksheet.cell(total_row, column, value)
        style_cell(cell, fill=PALE_GREEN, bold=True, color=NAVY)
    worksheet.cell(total_row, 5).number_format = "0.00%"
    worksheet.cell(total_row, 6).number_format = "0.00%"

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(17, column, header)
        style_cell(cell, fill=BLUE, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    combined_male = asn_gender["Pria"] + out_gender["Pria"]
    combined_female = asn_gender["Wanita"] + out_gender["Wanita"]
    combined_total = len(asn) + len(outsourcing)
    extra_rows = [
        ("Outsourcing", out_gender["Pria"], out_gender["Wanita"], len(outsourcing)),
        ("TOTAL SELURUH DATA", combined_male, combined_female, combined_total),
    ]
    for offset, (label, male, female, total) in enumerate(extra_rows, start=1):
        row = 17 + offset
        values = [label, male, female, total, male / total if total else 0, female / total if total else 0]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row, column, value)
            style_cell(cell, fill=PALE_GREEN if label.startswith("TOTAL") else PALE_BLUE, bold=label.startswith("TOTAL"), color=NAVY if label.startswith("TOTAL") else None)
        worksheet.cell(row, 5).number_format = "0.00%"
        worksheet.cell(row, 6).number_format = "0.00%"

    worksheet.merge_cells("A22:F22")
    worksheet["A22"] = (
        "Catatan: rekap utama adalah 843 ASN. Outsourcing ditampilkan terpisah agar tidak tercampur dengan status ASN. "
        "Semua 843 ASN dan 38 outsourcing memiliki data jenis kelamin yang dapat dinormalisasi."
    )
    worksheet["A22"].fill = PatternFill("solid", fgColor=YELLOW)
    worksheet["A22"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[22].height = 45

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "ASN Pria dan Wanita per Status"
    chart.y_axis.title = "Jumlah Pegawai"
    chart.x_axis.title = "Status"
    data = Reference(worksheet, min_col=2, max_col=3, min_row=9, max_row=13)
    categories = Reference(worksheet, min_col=1, min_row=10, max_row=13)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8.2
    chart.width = 16
    chart.legend.position = "b"
    worksheet.add_chart(chart, "H3")

    widths = {"A": 30, "B": 16, "C": 16, "D": 16, "E": 15, "F": 15, "G": 3}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A10"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1


def add_people_sheet(workbook: Workbook, title: str, people: list[dict[str, object]], table_name: str) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"] = "Jumlah pegawai"
    worksheet["B1"] = len(people)
    worksheet["A1"].font = Font(bold=True, color=NAVY)
    worksheet["B1"].font = Font(bold=True, color=NAVY, size=14)

    headers = ["No.", "Nama Pegawai", "NIP/NIK", "Status", "Unit Kerja", "Jenis Kelamin", "Sheet Sumber", "Baris Sumber"]
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(4, column, header)
        style_cell(cell, fill=BLUE, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sorted_people = sorted(people, key=lambda person: (str(person["status"]), str(person["name"])))
    for number, person in enumerate(sorted_people, start=1):
        row = 4 + number
        values = [
            number,
            person["name"],
            person["id"],
            person["status"],
            person["unit"],
            person["gender"],
            person["source_sheet"],
            person["source_row"],
        ]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        worksheet.cell(row, 3).number_format = "@"

    last_row = 4 + len(sorted_people)
    table = Table(displayName=table_name, ref=f"A4:H{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    worksheet.add_table(table)
    worksheet.auto_filter.ref = f"A4:H{last_row}"
    worksheet.freeze_panes = "A5"
    widths = [7, 42, 22, 22, 65, 17, 35, 16]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "4:4"


def add_methodology(workbook: Workbook) -> None:
    worksheet = workbook.create_sheet("METODOLOGI")
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "METODOLOGI REKAP"
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet["A1"].font = Font(color=WHITE, bold=True, size=15)
    worksheet["A1"].alignment = Alignment(horizontal="center")

    entries = [
        ("Cakupan utama", "ASN pada sheet DUK ASN BAPENDA KOTA MEDAN dengan NIP valid dan status PNS, CPNS, atau PPPK."),
        ("PPPK penuh/paruh waktu", "Ditentukan dengan pencocokan NIP terhadap sheet P3K PENUH WAKTU SESUAI UNIT dan P3K PARUH WAKTU SESUAI UNIT."),
        ("Jenis kelamin ASN", "Nilai Laki-laki dinormalisasi menjadi Pria; Perempuan ditampilkan sebagai Wanita."),
        ("Outsourcing", "Dihitung dari sheet OUTSOURCING berdasarkan nama dan NIK yang terisi; nilai Pria/PRIA dinormalisasi menjadi Pria."),
        ("Kontrol kelengkapan", "843 ASN dan 38 outsourcing memiliki jenis kelamin yang dapat dikenali; tidak ada data kosong/tidak valid."),
        ("Privasi", "Daftar bukti hanya membawa nama, NIP/NIK, status, unit, jenis kelamin, dan lokasi baris sumber; kontak dan alamat tidak disalin."),
    ]
    worksheet["A3"] = "Aspek"
    worksheet["B3"] = "Penjelasan"
    for cell in worksheet[3]:
        style_cell(cell, fill=BLUE, bold=True, color=WHITE)
    for row, (aspect, explanation) in enumerate(entries, start=4):
        worksheet.cell(row, 1, aspect).font = Font(bold=True, color=NAVY)
        worksheet.cell(row, 2, explanation).alignment = Alignment(wrap_text=True, vertical="top")
        worksheet.row_dimensions[row].height = 42
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 110
    worksheet.freeze_panes = "A4"


def main() -> None:
    asn, outsourcing = load_data()
    asn_gender = Counter(str(person["gender"]) for person in asn)
    out_gender = Counter(str(person["gender"]) for person in outsourcing)
    if asn_gender != Counter({"Pria": 465, "Wanita": 378}):
        raise ValueError(f"Kontrol jenis kelamin ASN gagal: {asn_gender}")
    if out_gender != Counter({"Pria": 20, "Wanita": 18}):
        raise ValueError(f"Kontrol jenis kelamin outsourcing gagal: {out_gender}")

    workbook = Workbook()
    add_summary(workbook, asn, outsourcing)
    add_people_sheet(workbook, "DAFTAR ASN PRIA", [person for person in asn if person["gender"] == "Pria"], "TabelAsnPria")
    add_people_sheet(workbook, "DAFTAR ASN WANITA", [person for person in asn if person["gender"] == "Wanita"], "TabelAsnWanita")
    add_people_sheet(
        workbook,
        "OUTSOURCING PRIA",
        [person for person in outsourcing if person["gender"] == "Pria"],
        "TabelOutsourcingPria",
    )
    add_people_sheet(
        workbook,
        "OUTSOURCING WANITA",
        [person for person in outsourcing if person["gender"] == "Wanita"],
        "TabelOutsourcingWanita",
    )
    add_methodology(workbook)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"ASN: pria {asn_gender['Pria']}, wanita {asn_gender['Wanita']}, total {len(asn)}")
    print(f"Outsourcing: pria {out_gender['Pria']}, wanita {out_gender['Wanita']}, total {len(outsourcing)}")


if __name__ == "__main__":
    main()
