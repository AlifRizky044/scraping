#!/usr/bin/env python3
"""Create an audit workbook for the August DUK update."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import update_duk_excel_agustus as updater


BASE_DIR = Path(__file__).resolve().parent
SOURCE = Path("/Users/alif/Downloads/DUK BAPENDA KOTA MEDAN 2026 1 Juni 2026 (1).xlsx")
UPDATED = BASE_DIR / "DUK BAPENDA KOTA MEDAN 2026 AGUSTUS UPDATED.xlsx"
JSON_PATH = BASE_DIR / "DUKPNSFULLagustus.json"
OUTPUT = BASE_DIR / "LAPORAN PEMBARUAN DUK AGUSTUS 2026.xlsx"

SHEET_NAMES = (
    "DUK ASN BAPENDA KOTA MEDAN",
    "PNS SESUAI UNIT KERJA",
    "P3K PENUH WAKTU SESUAI UNIT",
    "P3K PARUH WAKTU SESUAI UNIT",
)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def digits(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def excel_nip(value: object) -> str:
    number = digits(value)
    return number


def employees_by_nip(ws) -> dict[str, tuple]:
    result = {}
    for row in ws.iter_rows(values_only=True):
        nip = digits(row[2] if len(row) > 2 else "")
        if len(nip) >= 15:
            result[nip] = row
    return result


def add_table_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in range(1, ws.max_column + 1):
        values = [clean(ws.cell(row, column).value) for row in range(1, ws.max_row + 1)]
        width = min(max(max((len(value) for value in values), default=0) + 2, 12), 58)
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, header in enumerate(headers, start=1):
        if "NIP" in header.upper():
            for row_number in range(2, ws.max_row + 1):
                ws.cell(row_number, column).number_format = "@"


def main() -> None:
    august = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    august_by_nip = {digits(item.get("nip")): item for item in august}

    source_book = load_workbook(SOURCE, read_only=True, data_only=True)
    updated_book = load_workbook(UPDATED, read_only=True, data_only=True)
    source_main = employees_by_nip(source_book["DUK ASN BAPENDA KOTA MEDAN"])
    updated_main = employees_by_nip(updated_book["DUK ASN BAPENDA KOTA MEDAN"])

    workbook = Workbook()
    workbook.remove(workbook.active)

    recap = updated_book["REKAPITULASI"]
    summary_rows = [
        ["File sumber", SOURCE.name],
        ["File hasil", UPDATED.name],
        ["Total pegawai sheet utama", len(updated_main)],
        ["Total PNS/CPNS", recap["C16"].value],
        ["Total PPPK penuh waktu", recap["D16"].value],
        ["Total PPPK paruh waktu", recap["E16"].value],
        ["Total ASN pada rekap", recap["F16"].value],
        ["Selisih rekap terhadap sheet utama", (recap["F16"].value or 0) - len(updated_main)],
        ["PNS Bidang 4", recap["C8"].value],
        ["PNS UPT 4", recap["C12"].value],
        ["Jumlah perubahan pangkat Agustus", len(updater.AUGUST_RANK_UPDATES)],
        ["Catatan", "Roster aktif mengecualikan ANNA SARI HASIBUAN dan ASWAN karena telah meninggal; keduanya dicatat pada PENYESUAIAN ROSTER."],
    ]
    add_table_sheet(workbook, "RINGKASAN", ["Keterangan", "Nilai"], summary_rows)

    new_rows = []
    achmad = august_by_nip[updater.ACHMAD_NIP]
    new_rows.append(
        [excel_nip(updater.ACHMAD_NIP), achmad.get("namaWithGelar"), "BIDANG 1", "Tepat di bawah NIKMAL FAUZI LUBIS, S.T.; format PNS putih"]
    )
    for nip in updater.BIDANG_4_NIPS:
        employee = august_by_nip[nip]
        new_rows.append([excel_nip(nip), employee.get("namaWithGelar"), "BIDANG 4", "Di bawah FITRI ARUNA SIREGAR, S.SI, M.SI"])
    irma = august_by_nip[updater.IRMA_NIP]
    new_rows.append([excel_nip(updater.IRMA_NIP), irma.get("namaWithGelar"), "UPT 4", "Di bawah FAISAL, S.E, M.M"])
    add_table_sheet(workbook, "PEGAWAI BARU", ["NIP", "Nama", "Penempatan", "Posisi"], new_rows)

    source_people = {}
    for sheet_name in SHEET_NAMES:
        source_people.update(employees_by_nip(source_book[sheet_name]))
    roster_rows = []
    for name, (old_nip, correct_nip) in updater.NIP_CORRECTIONS.items():
        roster_rows.append(
            ["KOREKSI NIP", excel_nip(old_nip), excel_nip(correct_nip), name, "NIP lama diganti dengan NIP Agustus"]
        )
    removed_unique = updater.REMOVE_FROM_MAIN | updater.REMOVE_FROM_PNS | updater.REMOVE_FROM_PPPK_PART
    for nip in sorted(removed_unique):
        source_row = source_people.get(nip, ())
        roster_rows.append(
            ["DIKELUARKAN", excel_nip(nip), "", source_row[1] if source_row else "", "Tidak ada pada roster Agustus"]
        )
    for nip in (updater.ANNA_NIP, updater.ASWAN_NIP):
        employee = august_by_nip[nip]
        roster_rows.append(
            ["DIKELUARKAN (MENINGGAL)", excel_nip(nip), "", employee.get("namaWithGelar"), "Tidak dihitung sebagai ASN aktif"]
        )
    add_table_sheet(
        workbook,
        "PENYESUAIAN ROSTER",
        ["Tindakan", "NIP Lama", "NIP Agustus", "Nama", "Alasan"],
        roster_rows,
    )

    rank_rows = []
    for nip, new_group in updater.AUGUST_RANK_UPDATES.items():
        old = source_main.get(nip, ())
        new = updated_main[nip]
        note = "Detail SK/TMT pangkat Agustus tersedia"
        if nip in updater.RANK_DETAIL_NOT_AVAILABLE:
            note = "Nomor SK, pejabat penandatangan, dan TMT pangkat Agustus belum tersedia; kolom dikosongkan"
        rank_rows.append(
            [
                excel_nip(nip), new[1],
                old[8] if old else "", old[9] if old else "",
                new[8], new_group, new[10], new[11], new[12], note,
            ]
        )
    add_table_sheet(
        workbook,
        "PANGKAT AGUSTUS",
        ["NIP", "Nama", "Pangkat Lama", "Gol. Lama", "Pangkat Agustus", "Gol. Agustus", "No. SK", "Pejabat", "TMT", "Catatan"],
        rank_rows,
    )

    unmatched_rows = []
    for sheet_name in SHEET_NAMES:
        for nip, row in employees_by_nip(updated_book[sheet_name]).items():
            if nip not in august_by_nip:
                unmatched_rows.append(
                    [sheet_name, excel_nip(nip), row[1], "Tidak ditemukan dengan NIP yang sama pada DUKPNSFULLagustus.json", "Baris dipertahankan tanpa perubahan"]
                )
    for nip, employee in august_by_nip.items():
        if nip not in updated_main and nip not in updater.DECEASED_NIPS:
            unmatched_rows.append(
                [
                    "DAFTAR AGUSTUS → SHEET UTAMA",
                    excel_nip(nip),
                    employee.get("namaWithGelar"),
                    "Ada pada DUKPNSFULLagustus.json tetapi NIP tersebut tidak ada pada sheet utama workbook sumber",
                    "Dilaporkan; tidak disisipkan karena berada di luar tujuh penambahan yang diminta",
                ]
            )
    add_table_sheet(
        workbook,
        "NIP TIDAK COCOK",
        ["Sheet", "NIP Workbook", "Nama", "Masalah", "Tindakan"],
        unmatched_rows,
    )

    incomplete_rows = []
    for nip in sorted(updater.RANK_DETAIL_NOT_AVAILABLE):
        employee = august_by_nip[nip]
        latest = updater.latest_list_item(employee.get("riwayatPangkat"))
        incomplete_rows.append(
            [
                "PANGKAT", excel_nip(nip), employee.get("namaWithGelar"),
                employee.get("golonganTerakhir"), latest.get("pangkat"),
                "Riwayat masih menunjukkan pangkat lama", "No. SK/pejabat/TMT pangkat baru dikosongkan",
            ]
        )

    for employee in august:
        if digits(employee.get("nip")) in updater.DECEASED_NIPS:
            continue
        if clean(employee.get("statusKepegawaian")).upper() not in {"PNS", "CPNS"}:
            continue
        latest = updater.latest_list_item(employee.get("riwayatJabatan"))
        current_job = employee.get("jabatanTerakhir") or latest.get("jabatan", "")
        current_unit = employee.get("skpd") or latest.get("skpd", "")
        current_key = updater.canonical_job(current_job)
        history_key = updater.canonical_job(latest.get("jabatan", ""))
        current_unit_key = updater.canonical_job(current_unit)
        history_unit_key = updater.canonical_job(latest.get("skpd", ""))
        title_is_current = bool(
            current_key and history_key and (current_key in history_key or history_key in current_key)
        )
        unit_is_current = bool(
            current_unit_key
            and history_unit_key
            and (current_unit_key in history_unit_key or history_unit_key in current_unit_key)
        )
        history_is_current = title_is_current and unit_is_current
        if not history_is_current:
            incomplete_rows.append(
                [
                    "JABATAN", excel_nip(employee.get("nip")), employee.get("namaWithGelar"),
                    f"{current_job} — {current_unit}",
                    f"{latest.get('jabatan', '')} — {latest.get('skpd', '')}",
                    "Jabatan dan/atau unit Agustus berbeda dari riwayat jabatan terakhir",
                    "Jabatan/unit diperbarui; SK/TMT/pejabat jabatan dikosongkan",
                ]
            )
    add_table_sheet(
        workbook,
        "DETAIL BELUM TERSEDIA",
        ["Jenis", "NIP", "Nama", "Data Agustus", "Riwayat Detail Terakhir", "Masalah", "Tindakan"],
        incomplete_rows,
    )

    recap_rows = []
    for start, end, category in ((4, 15, "UNIT"), (21, 25, "JABATAN"), (31, 45, "PANGKAT"), (51, 57, "PENDIDIKAN")):
        for row_number in range(start, end + 1):
            recap_rows.append(
                [category, recap.cell(row_number, 2).value, recap.cell(row_number, 3).value,
                 recap.cell(row_number, 4).value, recap.cell(row_number, 5).value, recap.cell(row_number, 6).value]
            )
    add_table_sheet(
        workbook,
        "REKAP FINAL",
        ["Kategori", "Uraian", "Jumlah/PNS", "Laki-laki/PPPK Penuh", "Perempuan/PPPK Paruh", "Total ASN"],
        recap_rows,
    )

    workbook.save(OUTPUT)
    source_book.close()
    updated_book.close()
    print(f"Output: {OUTPUT.name}")
    print(f"NIP tidak cocok (kemunculan per sheet): {len(unmatched_rows)}")
    print(f"Detail belum tersedia: {len(incomplete_rows)}")


if __name__ == "__main__":
    main()
