#!/usr/bin/env python3
"""Membuat daftar nama di balik rekap pendidikan, usia, dan masa kerja 2025."""

from __future__ import annotations

import json
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / ".codex_deps"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

import buat_rekap_pegawai_31_des_2025 as base  # noqa: E402


OUTPUT = ROOT / "Daftar_Nama_Rincian_Rekap_2025.xlsx"
REF_DATE = date(2025, 12, 31)

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="D9E2F3")


def txt(value: Any) -> str:
    return "" if value is None else str(value).strip()


def title(ws, heading: str, subtitle: str, end_col: int) -> None:
    from openpyxl.utils import get_column_letter

    end = get_column_letter(end_col)
    ws.merge_cells(f"A1:{end}1")
    ws["A1"] = heading
    ws["A1"].font = Font(size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f"A2:{end}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="44546A")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 34
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


def style_header(cell) -> None:
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=GRID)


def style_data(cell, wrap: bool = False) -> None:
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    cell.border = Border(bottom=GRID)


def write_grouped_sheet(
    ws,
    heading: str,
    subtitle: str,
    headers: list[str],
    groups: list[tuple[str, list[dict[str, Any]], str]],
    values: Callable[[dict[str, Any], int], list[Any]],
    widths: dict[str, float],
) -> dict[str, int]:
    from openpyxl.utils import get_column_letter

    title(ws, heading, subtitle, len(headers))
    row = 4
    anchors: dict[str, int] = {}
    for label, records, color in groups:
        anchors[label] = row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        group_cell = ws.cell(row, 1, f"{label} — {len(records)} orang")
        group_cell.font = Font(size=12, bold=True, color="1F1F1F")
        group_cell.fill = PatternFill("solid", fgColor=color)
        group_cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 22
        row += 1
        for col, header in enumerate(headers, 1):
            style_header(ws.cell(row, col, header))
        row += 1
        if not records:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            cell = ws.cell(row, 1, "Tidak ada pegawai pada kategori ini.")
            cell.font = Font(italic=True, color="666666")
            cell.alignment = Alignment(vertical="center")
            row += 1
        else:
            for number, person in enumerate(records, 1):
                row_values = values(person, number)
                for col, value in enumerate(row_values, 1):
                    cell = ws.cell(row, col, value)
                    style_data(cell, wrap=col in {4, 5, len(headers)})
                    if headers[col - 1] in {"NIP", "NIK"}:
                        cell.number_format = "@"
                    if headers[col - 1] == "Tanggal Lahir" and value:
                        cell.number_format = "dd-mm-yyyy"
                row += 1
        row += 2

    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Rincian Rekap Pegawai — 31 Desember 2025"
    ws.oddFooter.right.text = "Halaman &P dari &N"
    return anchors


def make_roster() -> list[dict[str, Any]]:
    excel = base.extract_excel_sources(base.EXCEL_SOURCE)
    with base.JSON_SOURCE.open(encoding="utf-8") as fh:
        json_rows = json.load(fh)
    roster = base.enrich_roster(base.build_roster(json_rows, excel), excel)
    assert len(roster) == 848
    for person in roster:
        nip = txt(person.get("nip"))
        status = txt(person.get("statusKepegawaian"))
        start = base.start_from_nip(nip, status)
        tenure = base.completed_years(start, REF_DATE)
        person["_nip_start"] = start
        person["_nip_tenure"] = tenure
        if status == "PPPK":
            person["_nip_start_text"] = f"{nip[8:12]} (tahun saja)"
        elif start:
            person["_nip_start_text"] = start.strftime("%m-%Y")
        else:
            person["_nip_start_text"] = "Tidak terbaca"
        if tenure is None:
            person["_nip_tenure_group"] = "Tidak diketahui"
        elif tenure <= 10:
            person["_nip_tenure_group"] = "0–10 tahun"
        elif tenure <= 20:
            person["_nip_tenure_group"] = "11–20 tahun"
        elif tenure <= 30:
            person["_nip_tenure_group"] = "21–30 tahun"
        else:
            person["_nip_tenure_group"] = ">30 tahun"

        age = person.get("_age")
        if age is None:
            person["_age_group"] = "Tidak diketahui"
        elif age <= 30:
            person["_age_group"] = "≤30 tahun"
        elif age <= 45:
            person["_age_group"] = "31–45 tahun"
        else:
            person["_age_group"] = ">45 tahun"
    return roster


def group_records(roster: list[dict[str, Any]], key: str) -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for person in roster:
        groups[txt(person.get(key))].append(person)
    for records in groups.values():
        records.sort(key=lambda x: txt(x.get("namaWithGelar") or x.get("nama")))
    return groups


def main() -> None:
    roster = make_roster()
    education = group_records(roster, "_edu_label")
    ages = group_records(roster, "_age_group")
    tenures = group_records(roster, "_nip_tenure_group")

    wb = Workbook()
    summary = wb.active
    summary.title = "RINGKASAN"
    title(
        summary,
        "DAFTAR NAMA DI BALIK REKAP PEGAWAI 2025",
        "Posisi 31 Desember 2025; mencakup 848 ASN. Masa kerja dihitung hanya dari NIP sesuai keputusan terakhir.",
        6,
    )
    summary_headers = ["No.", "Dimensi", "Kategori", "Jumlah", "Daftar Nama", "Catatan"]
    for col, header in enumerate(summary_headers, 1):
        style_header(summary.cell(4, col, header))

    edu_order = ["S2", "D4/S1", "D3", "D1", "SMA/sederajat"]
    age_order = ["≤30 tahun", "31–45 tahun", ">45 tahun"]
    tenure_order = ["0–10 tahun", "11–20 tahun", "21–30 tahun", ">30 tahun"]

    summary_rows: list[tuple[str, str, int, str, str]] = []
    for category in edu_order:
        summary_rows.append(("Pendidikan", category, len(education.get(category, [])), "DAFTAR PENDIDIKAN", "Pendidikan tertinggi ASN"))
    other_education = sum(len(education.get(k, [])) for k in ("S3", "SD", "SMP/sederajat", "D2", "Tidak diketahui"))
    summary_rows.append(("Pendidikan", "Lainnya (S3/SD/SMP/D2)", other_education, "DAFTAR PENDIDIKAN", "Di luar kategori pada foto"))
    for category in age_order:
        summary_rows.append(("Usia", category, len(ages.get(category, [])), "DAFTAR USIA", "Usia tepat pada 31-12-2025"))
    for category in tenure_order:
        summary_rows.append(("Masa kerja", category, len(tenures.get(category, [])), "DAFTAR MASA KERJA NIP", "Tahun pengangkatan pada NIP"))

    for index, (dimension, category, count, sheet_name, note) in enumerate(summary_rows, 1):
        row = 4 + index
        values = [index, dimension, category, count, "Buka daftar", note]
        for col, value in enumerate(values, 1):
            cell = summary.cell(row, col, value)
            style_data(cell, wrap=col in {3, 6})
        link = summary.cell(row, 5)
        link.hyperlink = f"#'{sheet_name}'!A1"
        link.style = "Hyperlink"
        summary.cell(row, 4).alignment = Alignment(horizontal="center")

    summary.column_dimensions["A"].width = 7
    summary.column_dimensions["B"].width = 20
    summary.column_dimensions["C"].width = 34
    summary.column_dimensions["D"].width = 12
    summary.column_dimensions["E"].width = 18
    summary.column_dimensions["F"].width = 48
    summary.freeze_panes = "A5"
    summary.auto_filter.ref = f"A4:F{summary.max_row}"
    summary.page_setup.orientation = "landscape"
    summary.sheet_properties.pageSetUpPr.fitToPage = True

    # Pendidikan.
    ws_edu = wb.create_sheet("DAFTAR PENDIDIKAN")
    edu_groups = [
        ("S2", education.get("S2", []), LIGHT_GREEN),
        ("D4/S1", education.get("D4/S1", []), LIGHT_BLUE),
        ("D3", education.get("D3", []), LIGHT_ORANGE),
        ("D1", education.get("D1", []), GRAY),
        ("SMA/sederajat", education.get("SMA/sederajat", []), "FFF2CC"),
    ]
    extras = []
    for key in ("S3", "SD", "SMP/sederajat", "D2", "Tidak diketahui"):
        extras.extend(education.get(key, []))
    extras.sort(key=lambda x: txt(x.get("namaWithGelar") or x.get("nama")))
    edu_groups.append(("Lainnya — di luar parameter foto", extras, GRAY))

    def edu_values(person: dict[str, Any], number: int) -> list[Any]:
        return [
            number,
            txt(person.get("nip")),
            txt(person.get("namaWithGelar") or person.get("nama")),
            txt(person.get("statusKepegawaian")),
            txt(person.get("_edu_label")),
            txt(person.get("_edu_detail")),
            txt(person.get("skpd")),
        ]

    write_grouped_sheet(
        ws_edu,
        "DAFTAR NAMA MENURUT PENDIDIKAN TERTINGGI",
        "Kelompok pada foto: S2, D4/S1, D3, D1, dan SMA. Lima ASN di luar kelompok tersebut dicantumkan pada bagian terakhir.",
        ["No.", "NIP", "Nama", "Status", "Kelompok Pendidikan", "Detail Pendidikan", "Unit Kerja"],
        edu_groups,
        edu_values,
        {"A": 7, "B": 21, "C": 36, "D": 10, "E": 22, "F": 38, "G": 52},
    )

    # Usia.
    ws_age = wb.create_sheet("DAFTAR USIA")
    age_groups = [
        ("≤30 tahun", ages.get("≤30 tahun", []), LIGHT_GREEN),
        ("31–45 tahun", ages.get("31–45 tahun", []), LIGHT_BLUE),
        (">45 tahun", ages.get(">45 tahun", []), LIGHT_ORANGE),
    ]

    def age_values(person: dict[str, Any], number: int) -> list[Any]:
        return [
            number,
            txt(person.get("nip")),
            txt(person.get("namaWithGelar") or person.get("nama")),
            txt(person.get("statusKepegawaian")),
            person.get("_dob"),
            person.get("_age"),
            txt(person.get("_age_group")),
            txt(person.get("skpd")),
        ]

    write_grouped_sheet(
        ws_age,
        "DAFTAR NAMA MENURUT KELOMPOK USIA",
        "Usia dihitung tepat pada 31 Desember 2025. Tanggal lahir berasal dari JSON dengan fallback delapan digit awal NIP.",
        ["No.", "NIP", "Nama", "Status", "Tanggal Lahir", "Usia", "Kelompok Usia", "Unit Kerja"],
        age_groups,
        age_values,
        {"A": 7, "B": 21, "C": 36, "D": 10, "E": 15, "F": 9, "G": 18, "H": 52},
    )

    # Masa kerja berbasis NIP saja.
    ws_tenure = wb.create_sheet("DAFTAR MASA KERJA NIP")
    tenure_groups = [
        ("0–10 tahun", tenures.get("0–10 tahun", []), LIGHT_GREEN),
        ("11–20 tahun", tenures.get("11–20 tahun", []), LIGHT_BLUE),
        ("21–30 tahun", tenures.get("21–30 tahun", []), LIGHT_ORANGE),
        (">30 tahun", tenures.get(">30 tahun", []), "FFF2CC"),
    ]

    def tenure_values(person: dict[str, Any], number: int) -> list[Any]:
        return [
            number,
            txt(person.get("nip")),
            txt(person.get("namaWithGelar") or person.get("nama")),
            txt(person.get("statusKepegawaian")),
            txt(person.get("_nip_start_text")),
            person.get("_nip_tenure"),
            txt(person.get("_nip_tenure_group")),
            txt(person.get("skpd")),
        ]

    write_grouped_sheet(
        ws_tenure,
        "DAFTAR NAMA MENURUT MASA KERJA — BERDASARKAN NIP",
        "PNS/CPNS memakai tahun-bulan pengangkatan pada NIP. PPPK memakai tahun 2025; kode 21/22 sesudah tahun bukan bulan.",
        ["No.", "NIP", "Nama", "Status", "TMT dari NIP", "Masa Kerja (tahun)", "Kelompok Masa Kerja", "Unit Kerja"],
        tenure_groups,
        tenure_values,
        {"A": 7, "B": 21, "C": 36, "D": 10, "E": 19, "F": 20, "G": 24, "H": 52},
    )

    # Catatan ringkas.
    notes = wb.create_sheet("CATATAN")
    title(notes, "CATATAN PERHITUNGAN", "Definisi yang dipakai pada workbook daftar nama.", 3)
    note_headers = ["Topik", "Metode", "Catatan"]
    for col, header in enumerate(note_headers, 1):
        style_header(notes.cell(4, col, header))
    note_rows = [
        ("Cakupan", "848 ASN per 31-12-2025", "38 outsourcing tidak masuk daftar karena tidak memiliki NIP dan hanya tersedia sebagai agregat."),
        ("Pendidikan", "Pendidikan tertinggi", "Riwayat pendidikan JSON dengan fallback ke DUK Excel berdasarkan NIP/NIK."),
        ("Usia", "Tahun penuh pada 31-12-2025", "Kelompok: ≤30, 31–45, dan >45 tahun."),
        ("Masa kerja PNS/CPNS", "Enam digit tahun-bulan pada posisi 9–14 NIP", "Contoh 201101 berarti TMT Januari 2011."),
        ("Masa kerja PPPK", "Empat digit tahun pada posisi 9–12 NIP", "Kode 21/22 sesudah tahun bukan bulan; seluruh PPPK sumber bertahun pengangkatan 2025."),
        ("Koreksi tanggal efektif", "Sama dengan workbook rekap utama", "Tiga mutasi 2026 dikeluarkan; pegawai dengan BUP 01-01-2026 tetap dimasukkan."),
        ("Privasi", "Data lokal", "Workbook memuat NIP dan harus dibagikan secara terbatas."),
    ]
    for row_no, values in enumerate(note_rows, 5):
        for col, value in enumerate(values, 1):
            cell = notes.cell(row_no, col, value)
            style_data(cell, wrap=True)
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 48
    notes.column_dimensions["C"].width = 80
    notes.freeze_panes = "A5"

    wb.properties.title = "Daftar Nama Rincian Rekap Pegawai 2025"
    wb.properties.subject = "Rincian pendidikan, usia, dan masa kerja berbasis NIP"
    wb.properties.creator = "Codex"
    wb.save(OUTPUT)

    checks = {
        "output": str(OUTPUT),
        "total_asn": len(roster),
        "pendidikan": {k: len(education.get(k, [])) for k in edu_order},
        "pendidikan_lainnya": other_education,
        "usia": {k: len(ages.get(k, [])) for k in age_order},
        "masa_kerja_nip": {k: len(tenures.get(k, [])) for k in tenure_order},
    }
    print(json.dumps(checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
