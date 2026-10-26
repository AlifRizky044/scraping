from __future__ import annotations

import json
import re
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE_XLSX = Path("pk pw (1).xlsx")
SOURCE_JSON = Path("DUKPNSFULLagustus.json")
OUTPUT_XLSX = Path("pk pw (1) - dirapikan.xlsx")


MONTHS = {
    1: "JANUARI",
    2: "FEBRUARI",
    3: "MARET",
    4: "APRIL",
    5: "MEI",
    6: "JUNI",
    7: "JULI",
    8: "AGUSTUS",
    9: "SEPTEMBER",
    10: "OKTOBER",
    11: "NOVEMBER",
    12: "DESEMBER",
}

MONTH_PATTERN = "|".join(MONTHS.values())

DEGREE_MAP = {
    "S.E": "S.E.",
    "S.E.": "S.E.",
    "S.H": "S.H.",
    "S.KOM": "S.Kom.",
    "S.KOM.": "S.Kom.",
    "S.T": "S.T.",
    "A.MD": "A.Md.",
    "S.SOS": "S.Sos.",
    "S.M": "S.M.",
    "S.M.": "S.M.",
    "S.PD": "S.Pd.",
    "S.AK": "S.Ak.",
    "S.AK.": "S.Ak.",
    "SST": "S.ST.",
    "A.MD.KOM": "A.Md.Kom.",
    "S.SI": "S.Si.",
    "S.PSI": "S.Psi.",
    "S.I.KOM": "S.I.Kom.",
    "S.P": "S.P.",
    "A.MD.PJK": "A.Md.Pjk.",
    "A.MD.PJK.": "A.Md.Pjk.",
    "A.MD.AKUN": "A.Md.Akun.",
}


def collapse_spaces(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def digits(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def name_key(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", collapse_spaces(value).upper())


def format_degree(value: object) -> str:
    degree = collapse_spaces(value)
    if not degree:
        return ""
    key = degree.upper().rstrip(".")
    return DEGREE_MAP.get(key, DEGREE_MAP.get(degree.upper(), degree))


def format_person_name(profile: dict) -> tuple[str, str]:
    bare_name = collapse_spaces(profile.get("nama")).upper()
    prefix = collapse_spaces(profile.get("gelarDepan"))
    suffix = format_degree(profile.get("gelarBelakang"))

    full_name = bare_name
    if prefix:
        full_name = f"{prefix} {full_name}"
    if suffix:
        full_name = f"{full_name}, {suffix}"
    return full_name, bare_name


def extract_birthplace(value: object, fallback: object) -> str:
    text = collapse_spaces(value).upper()
    date_tail = re.compile(
        rf"\s*[/,]?\s*\d{{1,2}}\s*(?:[-/.]\s*|\s+)"
        rf"(?:{MONTH_PATTERN}|\d{{1,2}})\s*(?:[-/.]\s*|\s+)\d{{4}}\s*$"
    )
    place = date_tail.sub("", text).strip(" /,-")
    if not place or any(char.isdigit() for char in place):
        place = collapse_spaces(fallback).upper()
    return place


def format_birth(value: object, profile: dict) -> str:
    place = extract_birthplace(value, profile.get("tempatLahir"))
    birth_date = datetime.strptime(profile["tanggalLahir"], "%Y-%m-%d")
    return f"{place} / {birth_date.day:02d} {MONTHS[birth_date.month]} {birth_date.year}"


def education_major(profile: dict) -> str:
    history = profile.get("riwayatPendidikan")
    if not isinstance(history, list):
        return ""

    bachelors = []
    for entry in history:
        if not isinstance(entry, dict):
            continue
        level = collapse_spaces(entry.get("tingkatPendidikan")).upper()
        detail = collapse_spaces(entry.get("detailPendidikan")).upper()
        normalized = re.sub(r"[^A-Z0-9]", "", level + detail)
        if "S1" in normalized or "SARJANA" in normalized:
            bachelors.append(entry)

    if not bachelors:
        return ""

    latest = bachelors[-1]
    major = collapse_spaces(latest.get("jurusan")).upper()
    if major and major not in {"-", "SARJANA", "S-1", "S1"}:
        return major

    detail = collapse_spaces(latest.get("detailPendidikan")).upper()
    detail = re.sub(r"^S\s*[.-]?\s*1\s*(?:/\s*SARJANA)?\s*", "", detail)
    return detail.strip(" -/.,")


def format_education(value: object, profile: dict) -> str:
    text = collapse_spaces(value).upper()
    text = re.sub(r"^SARJANA\s+S\s*[-.]?\s*1\b\s*[-:]?\s*", "S-1 ", text)
    text = re.sub(r"^S\s*[-.]?\s*1\b\s*[-:]?\s*", "S-1 ", text)
    text = re.sub(r"^D\s*[-.]?\s*(?:III|3)\b\s*[-:]?\s*", "D-III ", text)
    text = re.sub(r"^D\s*[-.]?\s*(?:IV|4)\b\s*[-:]?\s*", "D-IV ", text)

    corrections = {
        r"\bAKUTANSI\b": "AKUNTANSI",
        r"\bAKUNTASI\b": "AKUNTANSI",
        r"\bADMINSTRASI\b": "ADMINISTRASI",
        r"\bTEKINIK\b": "TEKNIK",
        r"\bALIYA\b": "ALIYAH",
        r"\bMANAGEMENT\b": "MANAJEMEN",
        r"\bILMU\s+ILMU\b": "ILMU-ILMU",
    }
    for pattern, replacement in corrections.items():
        text = re.sub(pattern, replacement, text)

    text = re.sub(r"\s*,\s*", ", ", text)
    text = re.sub(r"\(\s+", "(", text)
    text = re.sub(r"\s+\)", ")", text)
    text = collapse_spaces(text).strip(" -")

    if text == "S-1":
        major = education_major(profile)
        if major:
            text = f"S-1 {major}"
    elif text == "S-1 KOM":
        major = education_major(profile)
        if major:
            text = f"S-1 {major}"

    return text


def format_address(value: object) -> str:
    text = collapse_spaces(value).upper()
    text = re.sub(r"\s*,\s*", ", ", text)
    return text


def format_unit(value: object, profile: dict) -> str:
    original = collapse_spaces(value).upper()
    profile_unit = collapse_spaces(profile.get("skpd")).upper()

    if "UPT" in original and profile_unit.startswith("UPT PAJAK DAN RETRIBUSI WILAYAH "):
        if "TATA USAHA" in original:
            return f"SUB BAGIAN TATA USAHA {profile_unit}"
        return profile_unit

    text = original
    corrections = {
        r"\bPENDAPTAN\b": "PENDAPATAN",
        r"\bPENDAPATN\b": "PENDAPATAN",
        r"\bDARRAH\b": "DAERAH",
        r"\bKOTA\s+MEDA\b": "KOTA MEDAN",
        r"\bSUB BIDANG TEKNIK PARKIR\b": "SUB BIDANG TEKNIS PARKIR",
        r"\bBEA PEROLEHAN ATAS TANAH\b": "BEA PEROLEHAN HAK ATAS TANAH",
        r"\bBADAN PENDAPATAN KOTA MEDAN\b": "BADAN PENDAPATAN DAERAH KOTA MEDAN",
    }
    for pattern, replacement in corrections.items():
        text = re.sub(pattern, replacement, text)
    text = re.sub(r"\s*,\s*", ", ", text)
    return collapse_spaces(text)


def copy_cell_format(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def main() -> None:
    with SOURCE_JSON.open(encoding="utf-8") as handle:
        profiles = json.load(handle)

    profiles_by_nip = {
        digits(profile.get("nip")): profile
        for profile in profiles
        if digits(profile.get("nip"))
    }
    profiles_by_name: dict[str, list[dict]] = {}
    for profile in profiles:
        profiles_by_name.setdefault(name_key(profile.get("nama")), []).append(profile)

    workbook = load_workbook(SOURCE_XLSX)
    sheet = workbook.active

    last_row_by_nip: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        last_row_by_nip[digits(sheet.cell(row, 3).value)] = row

    rows_to_delete = [
        row
        for row in range(2, sheet.max_row + 1)
        if row != last_row_by_nip[digits(sheet.cell(row, 3).value)]
    ]
    for row in sorted(rows_to_delete, reverse=True):
        sheet.delete_rows(row, 1)

    headers = [
        "NO",
        "NAMA DENGAN GELAR",
        "NIP",
        "TEMPAT / TANGGAL LAHIR",
        "PENDIDIKAN",
        "ALAMAT",
        "JABATAN",
        "UNIT KERJA",
        "NAMA TANPA GELAR",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column).value = header

    corrected_nips = 0
    matched_profiles = 0
    for row in range(2, sheet.max_row + 1):
        original_nip = digits(sheet.cell(row, 3).value)
        profile = profiles_by_nip.get(original_nip)

        if profile is None:
            candidates = profiles_by_name.get(name_key(sheet.cell(row, 2).value), [])
            if len(candidates) == 1:
                profile = candidates[0]
                corrected_nip = digits(profile.get("nip"))
                if corrected_nip and corrected_nip != original_nip:
                    sheet.cell(row, 3).value = corrected_nip
                    corrected_nips += 1

        if profile is None:
            raise ValueError(f"Profil JSON tidak ditemukan untuk baris {row}, NIP {original_nip}")

        matched_profiles += 1
        full_name, bare_name = format_person_name(profile)
        sheet.cell(row, 2).value = full_name
        sheet.cell(row, 3).value = digits(profile.get("nip"))
        sheet.cell(row, 3).number_format = "@"
        sheet.cell(row, 4).value = format_birth(sheet.cell(row, 4).value, profile)
        sheet.cell(row, 5).value = format_education(sheet.cell(row, 5).value, profile)
        sheet.cell(row, 6).value = format_address(sheet.cell(row, 6).value)
        sheet.cell(row, 7).value = collapse_spaces(profile.get("jabatanTerakhir")).upper()
        sheet.cell(row, 8).value = format_unit(sheet.cell(row, 8).value, profile)
        sheet.cell(row, 9).value = bare_name

        copy_cell_format(sheet.cell(row, 2), sheet.cell(row, 9))

    copy_cell_format(sheet.cell(1, 8), sheet.cell(1, 9))
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 23
    sheet.column_dimensions["D"].width = 37
    sheet.column_dimensions["E"].width = 38
    sheet.column_dimensions["F"].width = 62
    sheet.column_dimensions["G"].width = 36
    sheet.column_dimensions["H"].width = 80
    sheet.column_dimensions["I"].width = 38

    for table in sheet.tables.values():
        table.ref = f"A1:I{sheet.max_row}"

    workbook.save(OUTPUT_XLSX)
    print(f"Output: {OUTPUT_XLSX}")
    print(f"Baris data awal: {111 - 1}")
    print(f"Duplikat dibuang: {len(rows_to_delete)}")
    print(f"Baris data akhir: {sheet.max_row - 1}")
    print(f"Profil JSON cocok: {matched_profiles}")
    print(f"NIP dikoreksi dari kecocokan nama: {corrected_nips}")


if __name__ == "__main__":
    main()
