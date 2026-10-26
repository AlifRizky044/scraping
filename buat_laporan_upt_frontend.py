import os
import re
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import buat_laporan_upt_harian as base


BASE_DIR = Path(__file__).resolve().parent
BASE_XLSX = BASE_DIR / "LAPORAN UPT HARIAN.xlsx"
OUTPUT_XLSM = BASE_DIR / "LAPORAN UPT FRONTEND.xlsm"
VBA_DIR = BASE_DIR / "vba"
OVBA_REPO = BASE_DIR / ".codex_deps" / "ovba-writer"
OVBA_SEED = OVBA_REPO / "vbaproject" / "testdata" / "corpus" / "p1_compiled.bin"

DISTRICT_ROWS = [
    9, 10, 11,
    13, 14, 15,
    17, 18, 19,
    21, 22, 23,
    25, 26, 27,
    29, 30, 31,
    33, 34, 35,
]


def customize_frontend(source_path: Path, output_path: Path) -> None:
    wb = load_workbook(source_path, data_only=False)
    report_ws = wb["LAPORAN"]

    # Keep one worksheet so its VBA document module maps exactly to Sheet1.
    for worksheet in list(wb.worksheets):
        if worksheet is not report_ws:
            wb.remove(worksheet)
    for defined_name in list(wb.defined_names):
        del wb.defined_names[defined_name]

    report_ws["B3"] = (
        '=DATE($G$2,MATCH($E$3,$I$2:$I$13,0),'
        'MIN($E$2,DAY(EOMONTH(DATE($G$2,MATCH($E$3,$I$2:$I$13,0),1),0))))'
    )
    report_ws["G3"] = '=COUNT($L$40:$L$10000)'

    if "A5:G5" not in [str(item) for item in report_ws.merged_cells.ranges]:
        report_ws.merge_cells("A5:G5")
    report_ws["A5"] = (
        "INPUT LANGSUNG: pilih mode Harian dan tanggal, lalu ketik pada sel kuning. "
        "Aktifkan Macro / Enable Content agar data tersimpan per tanggal."
    )
    report_ws["A5"].fill = PatternFill("solid", fgColor="FFF2CC")
    report_ws["A5"].font = Font(name="Calibri", size=10, bold=True, color="9C5700")
    report_ws["A5"].alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    report_ws.row_dimensions[5].height = 32

    input_comment = Comment(
        "Ketik langsung pada mode Harian. Nilai disimpan otomatis menurut Tanggal Aktif.",
        "Laporan UPT",
    )
    for row in DISTRICT_ROWS:
        for col in (2, 3, 4, 5, 7):
            cell = report_ws.cell(row, col)
            cell.value = None
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            if row == DISTRICT_ROWS[0] and col == 2:
                cell.comment = input_comment
        report_ws.cell(row, 6, f"=IFERROR(D{row}/C{row},0)")
        report_ws.cell(row, 6).fill = PatternFill("solid", fgColor="E7E6E6")

    report_ws["H1"] = "Tanggal"
    for day in range(1, 32):
        report_ws.cell(day + 1, 8, day)
    report_ws["I1"] = "Bulan"
    for row, month in enumerate(base.MONTHS, 2):
        report_ws.cell(row, 9, month)
    report_ws["J1"] = "Tahun"
    for row, year in enumerate(range(2020, 2041), 2):
        report_ws.cell(row, 10, year)

    storage_headers = [
        "Tanggal", "Hari", "UPT", "Kecamatan", "Jumlah Petugas",
        "Target WP", "Realisasi SPPT", "Realisasi RP", "Keterangan",
    ]
    for column, value in enumerate(storage_headers, 12):
        report_ws.cell(39, column, value)
    report_ws["V1"] = "Kontrol Macro"
    report_ws["V2"] = "=B3"
    report_ws["V2"].number_format = "dd mmm yyyy"
    report_ws["V3"] = "Harian"

    report_ws.data_validations.dataValidation = []
    validations = [
        ("B2", DataValidation(type="list", formula1='"Harian,Bulanan"')),
        (
            "E2",
            DataValidation(
                type="list",
                formula1=(
                    "=OFFSET($H$2,0,0,DAY(EOMONTH(DATE($G$2,"
                    "MATCH($E$3,$I$2:$I$13,0),1),0)),1)"
                ),
            ),
        ),
        ("E3", DataValidation(type="list", formula1="=$I$2:$I$13")),
        ("G2", DataValidation(type="list", formula1="=$J$2:$J$22")),
    ]
    for coordinate, validation in validations:
        validation.showDropDown = False
        validation.showErrorMessage = True
        validation.errorTitle = "Pilihan tidak valid"
        validation.error = "Pilih nilai dari dropdown."
        report_ws.add_data_validation(validation)
        validation.add(coordinate)

    for column in range(8, 23):
        report_ws.column_dimensions[get_column_letter(column)].hidden = True

    report_ws.sheet_properties.codeName = "Sheet1"
    report_ws.sheet_properties.tabColor = "70AD47"
    wb.code_name = "ThisWorkbook"

    wb.active = 0
    report_ws.sheet_view.selection[0].activeCell = "B9"
    report_ws.sheet_view.selection[0].sqref = "B9"
    wb.save(output_path)
    wb.close()


def build_vba_project(temp_dir: Path) -> Path:
    result = temp_dir / "vbaProject.bin"
    env = os.environ.copy()
    env["GOCACHE"] = str(BASE_DIR / ".codex_deps" / "go-cache")
    env["GOPATH"] = str(BASE_DIR / ".codex_deps" / "go-path")
    command = [
        "go",
        "run",
        "./cmd/rewrite",
        str(OVBA_SEED),
        str(VBA_DIR / "ThisWorkbook.cls"),
        str(VBA_DIR / "ModuleUPT.bas"),
        str(result),
    ]
    subprocess.run(command, cwd=OVBA_REPO, env=env, check=True)
    if not result.exists() or result.stat().st_size == 0:
        raise RuntimeError("vbaProject.bin gagal dibuat")
    return result


def make_macro_enabled(source_xlsx: Path, vba_project: Path, output_xlsm: Path) -> None:
    with ZipFile(source_xlsx, "r") as source_zip, ZipFile(
        output_xlsm, "w", ZIP_DEFLATED
    ) as target_zip:
        rels_xml = source_zip.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        rel_ids = [int(value) for value in re.findall(r'Id="rId(\d+)"', rels_xml)]
        next_rel_id = max(rel_ids, default=0) + 1
        vba_relationship = (
            f'<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
            f'Target="vbaProject.bin" Id="rId{next_rel_id}"/>'
        )
        rels_xml = rels_xml.replace("</Relationships>", vba_relationship + "</Relationships>")

        content_types = source_zip.read("[Content_Types].xml").decode("utf-8")
        content_types = content_types.replace(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        )
        vba_override = (
            '<Override PartName="/xl/vbaProject.bin" '
            'ContentType="application/vnd.ms-office.vbaProject"/>'
        )
        content_types = content_types.replace("</Types>", vba_override + "</Types>")

        workbook_xml = source_zip.read("xl/workbook.xml").decode("utf-8")
        if "codeName=" not in workbook_xml:
            workbook_xml = workbook_xml.replace(
                "<workbookPr/>", '<workbookPr codeName="ThisWorkbook"/>'
            )

        replaced = {
            "[Content_Types].xml": content_types.encode("utf-8"),
            "xl/_rels/workbook.xml.rels": rels_xml.encode("utf-8"),
            "xl/workbook.xml": workbook_xml.encode("utf-8"),
        }
        for info in source_zip.infolist():
            if info.filename in replaced:
                target_zip.writestr(info, replaced[info.filename])
            else:
                target_zip.writestr(info, source_zip.read(info.filename))
        target_zip.write(vba_project, "xl/vbaProject.bin")


def build_frontend() -> Path:
    base.build_workbook()
    with tempfile.TemporaryDirectory(prefix="upt_frontend_") as temp_name:
        temp_dir = Path(temp_name)
        frontend_xlsx = temp_dir / "frontend.xlsx"
        customize_frontend(BASE_XLSX, frontend_xlsx)
        vba_project = build_vba_project(temp_dir)
        make_macro_enabled(frontend_xlsx, vba_project, OUTPUT_XLSM)
    return OUTPUT_XLSM


if __name__ == "__main__":
    result = build_frontend()
    print(result)
