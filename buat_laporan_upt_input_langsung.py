from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.filters import DateGroupItem, FilterColumn, Filters


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / "LAPORAN UPT - INPUT LANGSUNG.xlsx"

UPT_DATA = [
    ("UPT I", ["Medan Amplas", "Medan Area", "Medan Kota"]),
    ("UPT II", ["Medan Tembung", "Medan Perjuangan", "Medan Denai"]),
    ("UPT III", ["Medan Sunggal", "Medan Petisah", "Medan Helvetia"]),
    ("UPT IV", ["Medan Barat", "Medan Timur", "Medan Deli"]),
    ("UPT V", ["Medan Baru", "Medan Polonia", "Medan Maimun"]),
    ("UPT VI", ["Medan Johor", "Medan Selayang", "Medan Tuntungan"]),
    ("UPT VII", ["Medan Belawan", "Medan Marelan", "Medan Labuhan"]),
]

YEAR = date.today().year
START_DATE = date(YEAR, 1, 1)
END_DATE = date(YEAR, 12, 31)
TODAY = date.today()

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
WHITE = "FFFFFF"
GRAY = "666666"
THIN_SIDE = Side(style="thin", color="B7B7B7")
MEDIUM_SIDE = Side(style="medium", color=BLUE)


def all_dates(start, end):
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def box(side=THIN_SIDE):
    return Border(left=side, right=side, top=side, bottom=side)


def style_header(cell):
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = box()


def setup_input_sheet(ws):
    ws.title = "KECAMATAN"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C7"

    ws.merge_cells("A1:H1")
    ws["A1"] = "LAPORAN HARIAN UPT / KECAMATAN"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Calibri", size=17, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws["A2"] = "HARI INI"
    ws["B2"] = '=CHOOSE(WEEKDAY(TODAY()),"Minggu","Senin","Selasa","Rabu","Kamis","Jumat","Sabtu")'
    ws["C2"] = "TANGGAL HARI INI"
    ws["D2"] = "=TODAY()"
    ws["D2"].number_format = "dd mmmm yyyy"
    ws["E2"] = "PERIODE DATA"
    ws["F2"] = f"1 Januari - 31 Desember {YEAR}"
    ws.merge_cells("F2:H2")
    for cell in ws[2]:
        cell.border = box()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for coordinate in ("A2", "C2", "E2"):
        ws[coordinate].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws[coordinate].font = Font(bold=True, color=NAVY)
    for coordinate in ("B2", "D2", "F2"):
        ws[coordinate].fill = PatternFill("solid", fgColor=PALE_GREEN)
        ws[coordinate].font = Font(bold=True, color="375623")

    ws.merge_cells("A3:H3")
    ws["A3"] = (
        "INPUT LANGSUNG DI TABEL INI. KLIK DROPDOWN TANGGAL DI SEL A6, "
        "hapus centang Pilih Semua, lalu pilih tanggal 1 sampai akhir bulan."
    )
    ws["A3"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    ws["A3"].font = Font(bold=True, color="7F6000")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 30

    ws.merge_cells("E5:F5")
    ws["A5"] = "FILTER TANGGAL ↓"
    ws["B5"] = "Wilayah"
    ws["C5"] = "Petugas"
    ws["D5"] = "Target"
    ws["E5"] = "Realisasi"
    ws["G5"] = "Capaian"
    ws["H5"] = "Catatan"
    ws["A6"] = "Tanggal"
    ws["B6"] = "UPT / Kecamatan"
    ws["C6"] = "Jumlah Petugas"
    ws["D6"] = "Target WP"
    ws["E6"] = "Jumlah SPPT"
    ws["F6"] = "Jumlah RP"
    ws["G6"] = "Persentase"
    ws["H6"] = "Keterangan"
    for row in range(5, 7):
        for col in range(1, 9):
            style_header(ws.cell(row, col))
    ws.row_dimensions[5].height = 27
    ws.row_dimensions[6].height = 25

    current_row = 7
    source_links = []
    today_rows = []
    for work_date in all_dates(START_DATE, END_DATE):
        day_start = current_row
        upt_total_rows = []

        for upt, districts in UPT_DATA:
            upt_row = current_row
            upt_total_rows.append(upt_row)
            ws.cell(upt_row, 1, work_date)
            ws.cell(upt_row, 2, upt)
            child_rows = []

            for district in districts:
                current_row += 1
                row = current_row
                child_rows.append(row)
                ws.cell(row, 1, work_date)
                ws.cell(row, 2, district)
                for col in range(3, 7):
                    ws.cell(row, col, 0)
                ws.cell(row, 7, f"=IFERROR(E{row}/D{row},0)")
                ws.cell(row, 8, "")
                source_links.append((work_date, upt, district, row))

                for col in range(1, 9):
                    cell = ws.cell(row, col)
                    cell.border = box()
                    cell.alignment = Alignment(vertical="center", wrap_text=(col == 8))
                    if col in (3, 4, 5, 6, 8):
                        cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)

            for col in range(3, 7):
                letter = ws.cell(upt_row, col).column_letter
                ws.cell(upt_row, col, f"=SUM({letter}{child_rows[0]}:{letter}{child_rows[-1]})")
            ws.cell(upt_row, 7, f"=IFERROR(E{upt_row}/D{upt_row},0)")
            ws.cell(upt_row, 8, f"Total {upt}")
            for col in range(1, 9):
                cell = ws.cell(upt_row, col)
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
                cell.font = Font(bold=True, color=NAVY)
                cell.border = box(MEDIUM_SIDE)
                cell.alignment = Alignment(vertical="center", wrap_text=(col == 8))
            current_row += 1

        grand_row = current_row
        ws.cell(grand_row, 1, work_date)
        ws.cell(grand_row, 2, "GRAND TOTAL")
        for col in range(3, 7):
            letter = ws.cell(grand_row, col).column_letter
            refs = ",".join(f"{letter}{row}" for row in upt_total_rows)
            ws.cell(grand_row, col, f"=SUM({refs})")
        ws.cell(grand_row, 7, f"=IFERROR(E{grand_row}/D{grand_row},0)")
        ws.cell(grand_row, 8, "Total Harian")
        for col in range(1, 9):
            cell = ws.cell(grand_row, col)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color=WHITE)
            cell.border = box(MEDIUM_SIDE)
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 8))

        day_end = grand_row
        if work_date == TODAY:
            today_rows.extend(range(day_start, day_end + 1))
        else:
            for row in range(day_start, day_end + 1):
                ws.row_dimensions[row].hidden = True
        current_row += 1

    last_row = current_row - 1
    for row in range(7, last_row + 1):
        ws.cell(row, 1).number_format = "dd mmm yyyy"
        for col in (3, 4, 5):
            ws.cell(row, col).number_format = "#,##0"
        ws.cell(row, 6).number_format = '"Rp" #,##0'
        ws.cell(row, 7).number_format = "0.00%"

    ws.auto_filter.ref = f"A6:H{last_row}"
    ws.auto_filter.filterColumn.append(
        FilterColumn(
            colId=0,
            filters=Filters(
                dateGroupItem=[
                    DateGroupItem(
                        year=TODAY.year,
                        month=TODAY.month,
                        day=TODAY.day,
                        dateTimeGrouping="day",
                    )
                ]
            ),
        )
    )
    ws.sheet_properties.filterMode = True
    ws.auto_filter.add_sort_condition(f"A7:A{last_row}")

    ws["A6"].comment = Comment(
        "Klik panah filter pada sel ini. Hapus centang Pilih Semua, lalu pilih "
        "satu tanggal. Tabel akan menampilkan data milik tanggal tersebut.",
        "Laporan UPT",
    )

    ws.conditional_formatting.add(
        f"G7:G{last_row}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="F8696B",
            mid_type="num",
            mid_value=0.75,
            mid_color="FFEB84",
            end_type="num",
            end_value=1,
            end_color="63BE7B",
        ),
    )

    visible_input_row = today_rows[1] if today_rows else 8
    visible_input_cell = f"C{visible_input_row}"
    ws[visible_input_cell].comment = Comment(
        "Ketik angka langsung di sel kuning. Data tanggal lain tetap tersimpan; "
        "ganti tanggal melalui dropdown filter di kolom Tanggal.",
        "Laporan UPT",
    )
    widths = {
        "A": 19,
        "B": 25,
        "C": 16,
        "D": 15,
        "E": 17,
        "F": 21,
        "G": 14,
        "H": 35,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.topLeftCell = f"A{today_rows[0]}" if today_rows else "A7"
    ws.sheet_view.selection[0].activeCell = visible_input_cell
    ws.sheet_view.selection[0].sqref = visible_input_cell
    ws.sheet_properties.tabColor = "70AD47"
    return source_links, last_row, today_rows


def setup_source_sheet(ws, source_links):
    ws.title = "SUMBER DATA"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"

    ws.merge_cells("A1:H1")
    ws["A1"] = "SUMBER DATA OTOMATIS"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "Data pada sheet ini mengikuti input di KECAMATAN secara otomatis. "
        "Input tetap dilakukan di sheet KECAMATAN."
    )
    ws["A2"].fill = PatternFill("solid", fgColor="E7E6E6")
    ws["A2"].font = Font(italic=True, color=GRAY)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)

    headers = [
        "Tanggal",
        "UPT",
        "Kecamatan",
        "Jumlah Petugas",
        "Target WP",
        "Jumlah SPPT",
        "Jumlah RP",
        "Keterangan",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
        style_header(ws.cell(4, col))

    for source_row, (work_date, upt, district, input_row) in enumerate(source_links, 5):
        ws.cell(source_row, 1, work_date)
        ws.cell(source_row, 2, upt)
        ws.cell(source_row, 3, district)
        for source_col, input_col in zip(range(4, 8), (3, 4, 5, 6)):
            letter = ws.cell(input_row, input_col).column_letter
            ws.cell(
                source_row,
                source_col,
                f'=IF(KECAMATAN!{letter}{input_row}="","",KECAMATAN!{letter}{input_row})',
            )
        ws.cell(
            source_row,
            8,
            f'=IF(KECAMATAN!H{input_row}="","",KECAMATAN!H{input_row})',
        )
        for col in range(1, 9):
            cell = ws.cell(source_row, col)
            cell.border = box()
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 8))
        ws.cell(source_row, 1).number_format = "dd mmm yyyy"
        for col in (4, 5, 6):
            ws.cell(source_row, col).number_format = "#,##0"
        ws.cell(source_row, 7).number_format = '"Rp" #,##0'

    last_row = 4 + len(source_links)
    ws.auto_filter.ref = f"A4:H{last_row}"
    widths = {
        "A": 16,
        "B": 13,
        "C": 24,
        "D": 16,
        "E": 15,
        "F": 17,
        "G": 21,
        "H": 35,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.sheet_properties.tabColor = "A5A5A5"


def build_workbook():
    wb = Workbook()
    input_ws = wb.active
    source_ws = wb.create_sheet()
    source_links, _, _ = setup_input_sheet(input_ws)
    setup_source_sheet(source_ws, source_links)

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_workbook())
