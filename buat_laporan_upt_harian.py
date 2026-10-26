from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties


BASE_DIR = Path(__file__).resolve().parent
SOURCE = BASE_DIR / "LAPORAN UPT.xlsx"
OUTPUT = BASE_DIR / "LAPORAN UPT HARIAN.xlsx"
INPUT_FIRST_ROW = 5
INPUT_LAST_ROW = 10000


UPT_DATA = [
    ("UPT I", ["Medan Amplas", "Medan Area", "Medan Kota"]),
    ("UPT II", ["Medan Tembung", "Medan Perjuangan", "Medan Denai"]),
    ("UPT III", ["Medan Sunggal", "Medan Petisah", "Medan Helvetia"]),
    ("UPT IV", ["Medan Barat", "Medan Timur", "Medan Deli"]),
    ("UPT V", ["Medan Baru", "Medan Polonia", "Medan Maimun"]),
    ("UPT VI", ["Medan Johor", "Medan Selayang", "Medan Tuntungan"]),
    ("UPT VII", ["Medan Belawan", "Medan Marelan", "Medan Labuhan"]),
]

MONTHS = [
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]


NAVY = "17365D"
BLUE = "1F4E78"
TEAL = "0F6B78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
PALE_RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "666666"
LIGHT_GRAY = "E7E6E6"
MEDIUM_GRAY = "B7B7B7"
GREEN = "70AD47"
RED = "C00000"

THIN_GRAY = Side(style="thin", color="B7B7B7")
MEDIUM_BLUE = Side(style="medium", color=BLUE)


def title_style(cell, size=16):
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Calibri", size=size, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def header_style(cell):
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def add_defined_name(workbook, name, reference):
    workbook.defined_names.add(DefinedName(name, attr_text=reference))


def report_sum_formula(metric, row):
    metric_columns = {
        "Jumlah Petugas": "E",
        "Target WP": "F",
        "Realisasi SPPT": "G",
        "Realisasi RP": "H",
    }
    metric_col = metric_columns[metric]
    value_range = f"'INPUT HARIAN'!${metric_col}${INPUT_FIRST_ROW}:${metric_col}${INPUT_LAST_ROW}"
    district_range = f"'INPUT HARIAN'!$D${INPUT_FIRST_ROW}:$D${INPUT_LAST_ROW}"
    date_range = f"'INPUT HARIAN'!$A${INPUT_FIRST_ROW}:$A${INPUT_LAST_ROW}"
    return (
        f'=IF($B$3="","",IF($B$2="Harian",'
        f'IFERROR(SUMIFS({value_range},{district_range},$A{row},{date_range},$B$3),0),'
        f'IFERROR(SUMIFS({value_range},{district_range},$A{row},'
        f'{date_range},">="&EOMONTH($B$3,-1)+1,'
        f'{date_range},"<="&EOMONTH($B$3,0)),0)))'
    )


def report_note_formula(row):
    date_range = f"'INPUT HARIAN'!$A${INPUT_FIRST_ROW}:$A${INPUT_LAST_ROW}"
    district_range = f"'INPUT HARIAN'!$D${INPUT_FIRST_ROW}:$D${INPUT_LAST_ROW}"
    note_range = f"'INPUT HARIAN'!$I${INPUT_FIRST_ROW}:$I${INPUT_LAST_ROW}"
    return (
        f'=IF($B$3="","",IF($B$2="Harian",'
        f'IFERROR(LOOKUP(2,1/(({district_range}=$A{row})*'
        f'({date_range}=$B$3)*({note_range}<>"")),{note_range}),""),'
        f'IFERROR(LOOKUP(2,1/(({district_range}=$A{row})*'
        f'({date_range}>=EOMONTH($B$3,-1)+1)*'
        f'({date_range}<=EOMONTH($B$3,0))*'
        f'({note_range}<>"")),{note_range}),"")))'
    )


def setup_reference_sheet(ws):
    ws.title = "REFERENSI"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Kecamatan"
    ws["B1"] = "UPT"
    row = 2
    for upt, districts in UPT_DATA:
        for district in districts:
            ws.cell(row, 1, district)
            ws.cell(row, 2, upt)
            row += 1

    ws["D1"] = "Mode Laporan"
    ws["D2"] = "Harian"
    ws["D3"] = "Bulanan"

    ws["H1"] = "Tanggal"
    for day in range(1, 32):
        ws.cell(day + 1, 8, day)

    ws["I1"] = "Bulan"
    for row, month in enumerate(MONTHS, 2):
        ws.cell(row, 9, month)

    ws["J1"] = "Tahun"
    for row, year in enumerate(range(2020, 2041), 2):
        ws.cell(row, 10, year)

    for cell in ws[1]:
        if cell.value:
            header_style(cell)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["H"].width = 12
    ws.column_dimensions["I"].width = 16
    ws.column_dimensions["J"].width = 12


def setup_input_sheet(ws):
    ws.title = "INPUT HARIAN"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:I1")
    ws["A1"] = "INPUT DATA HARIAN UPT"
    title_style(ws["A1"], 15)
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    ws["A2"] = (
        "Isi kolom kuning. Satu baris = satu kecamatan pada satu tanggal; "
        "Hari dan UPT terisi otomatis."
    )
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color=GRAY)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 22

    headers = [
        "Tanggal",
        "Hari",
        "UPT",
        "Kecamatan",
        "Jumlah Petugas",
        "Target WP",
        "Realisasi SPPT",
        "Realisasi RP",
        "Keterangan",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
        header_style(ws.cell(4, col))
    ws.row_dimensions[4].height = 36

    first_data_row = INPUT_FIRST_ROW
    last_table_row = INPUT_LAST_ROW
    for row in range(first_data_row, last_table_row + 1):
        ws.cell(row, 2, f'=IF(A{row}="","",CHOOSE(WEEKDAY(A{row},2),"Senin","Selasa","Rabu","Kamis","Jumat","Sabtu","Minggu"))')
        ws.cell(row, 3, f'=IFERROR(VLOOKUP(D{row},\'REFERENSI\'!$A$2:$B$22,2,FALSE),"")')
        for col in range(1, 10):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
            )
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 9))
            if col in (1, 4, 5, 6, 7, 8, 9):
                cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
            else:
                cell.fill = PatternFill("solid", fgColor=PALE_BLUE)

        ws.cell(row, 1).number_format = "dd mmm yyyy"
        for col in (5, 6, 7):
            ws.cell(row, col).number_format = '#,##0'
        ws.cell(row, 8).number_format = '"Rp" #,##0'

    table = Table(displayName="tblInput", ref=f"A4:I{last_table_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    max_input_row = INPUT_LAST_ROW
    date_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2040,12,31)",
        allow_blank=True,
    )
    date_validation.promptTitle = "Tanggal input"
    date_validation.prompt = "Masukkan tanggal antara 1 Jan 2020 dan 31 Des 2040."
    date_validation.errorTitle = "Tanggal tidak valid"
    date_validation.error = "Gunakan tanggal yang valid (2020-2040)."
    date_validation.errorStyle = "stop"
    date_validation.showErrorMessage = True
    date_validation.showInputMessage = True
    ws.add_data_validation(date_validation)
    date_validation.add(f"A{first_data_row}:A{max_input_row}")

    district_validation = DataValidation(
        type="list", formula1="=DaftarKecamatan", allow_blank=True
    )
    district_validation.promptTitle = "Pilih kecamatan"
    district_validation.prompt = "Pilih kecamatan dari dropdown."
    district_validation.errorTitle = "Kecamatan tidak valid"
    district_validation.error = "Pilih nama kecamatan yang tersedia."
    district_validation.errorStyle = "stop"
    district_validation.showErrorMessage = True
    district_validation.showInputMessage = True
    ws.add_data_validation(district_validation)
    district_validation.add(f"D{first_data_row}:D{max_input_row}")

    integer_validation = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True
    )
    integer_validation.errorTitle = "Angka tidak valid"
    integer_validation.error = "Isi dengan bilangan bulat nol atau lebih."
    integer_validation.showErrorMessage = True
    ws.add_data_validation(integer_validation)
    integer_validation.add(f"E{first_data_row}:G{max_input_row}")

    money_validation = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True
    )
    money_validation.errorTitle = "Nilai tidak valid"
    money_validation.error = "Isi dengan angka nol atau lebih."
    money_validation.showErrorMessage = True
    ws.add_data_validation(money_validation)
    money_validation.add(f"H{first_data_row}:H{max_input_row}")

    duplicate_formula = (
        f'=AND($A{first_data_row}<>"",$D{first_data_row}<>"",'
        f'COUNTIFS($A${INPUT_FIRST_ROW}:$A${INPUT_LAST_ROW},$A{first_data_row},'
        f'$D${INPUT_FIRST_ROW}:$D${INPUT_LAST_ROW},$D{first_data_row})>1)'
    )
    ws.conditional_formatting.add(
        f"A{first_data_row}:I{max_input_row}",
        FormulaRule(
            formula=[duplicate_formula],
            fill=PatternFill("solid", fgColor=PALE_RED),
            font=Font(color=RED),
        ),
    )

    widths = {
        "A": 15,
        "B": 13,
        "C": 12,
        "D": 23,
        "E": 16,
        "F": 14,
        "G": 17,
        "H": 20,
        "I": 38,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.auto_filter.ref = f"A4:I{last_table_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    ws.print_area = "A1:I104"
    ws.oddFooter.center.text = "Halaman &P dari &N"


def setup_report_sheet(ws):
    ws.title = "LAPORAN"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    ws.merge_cells("A1:G1")
    ws["A1"] = "LAPORAN HARIAN / BULANAN UPT"
    title_style(ws["A1"], 17)
    ws.row_dimensions[1].height = 32

    ws["A2"] = "Mode Laporan"
    ws["A3"] = "Tanggal Aktif"
    ws["A4"] = "Periode Tampil"
    for row in range(2, 5):
        ws.cell(row, 1).font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row, 1).alignment = Alignment(vertical="center")
        ws.cell(row, 1).border = Border(
            left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
        )

    ws["B2"] = "Harian"
    month_match = "MATCH($E$3,'REFERENSI'!$I$2:$I$13,0)"
    ws["B3"] = (
        f'=DATE($G$2,{month_match},'
        f'MIN($E$2,DAY(EOMONTH(DATE($G$2,{month_match},1),0))))'
    )
    ws["B3"].number_format = "dd mmmm yyyy"
    months = 'CHOOSE(MONTH($B$3),"Januari","Februari","Maret","April","Mei","Juni","Juli","Agustus","September","Oktober","November","Desember")'
    days = 'CHOOSE(WEEKDAY($B$3),"Minggu","Senin","Selasa","Rabu","Kamis","Jumat","Sabtu")'
    ws["B4"] = (
        f'=IF($B$3="","Pilih tanggal",IF($B$2="Harian",'
        f'{days}&", "&TEXT($B$3,"dd")&" "&{months}&" "&YEAR($B$3),'
        f'{months}&" "&YEAR($B$3)))'
    )
    ws.merge_cells("B4:G4")

    for coordinate in ("B2",):
        cell = ws[coordinate]
        cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
        cell.font = Font(name="Calibri", size=11, bold=True, color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=MEDIUM_BLUE,
            right=MEDIUM_BLUE,
            top=MEDIUM_BLUE,
            bottom=MEDIUM_BLUE,
        )
    ws["B3"].fill = PatternFill("solid", fgColor=PALE_GREEN)
    ws["B3"].font = Font(name="Calibri", size=11, bold=True, color="375623")
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B3"].border = Border(
        left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
    )
    ws["B4"].fill = PatternFill("solid", fgColor=PALE_GREEN)
    ws["B4"].font = Font(name="Calibri", size=11, bold=True, color="375623")
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center")
    for row in range(2, 5):
        ws.row_dimensions[row].height = 24

    ws["D2"] = "Tanggal"
    ws["E2"] = "=DAY(TODAY())"
    ws["D3"] = "Bulan"
    ws["E3"] = '=CHOOSE(MONTH(TODAY()),"Januari","Februari","Maret","April","Mei","Juni","Juli","Agustus","September","Oktober","November","Desember")'
    for coordinate in ("D2", "D3"):
        ws[coordinate].font = Font(name="Calibri", size=9, bold=True, color=NAVY)
        ws[coordinate].alignment = Alignment(horizontal="right", vertical="center")
    for coordinate in ("E2", "E3"):
        ws[coordinate].fill = PatternFill("solid", fgColor=PALE_YELLOW)
        ws[coordinate].font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        ws[coordinate].alignment = Alignment(horizontal="center", vertical="center")
        ws[coordinate].border = Border(
            left=MEDIUM_BLUE,
            right=MEDIUM_BLUE,
            top=MEDIUM_BLUE,
            bottom=MEDIUM_BLUE,
        )

    ws["F2"] = "Tahun"
    ws["G2"] = "=YEAR(TODAY())"
    ws["F3"] = "Jumlah Entri"
    ws["G3"] = f'=COUNT(\'INPUT HARIAN\'!$A${INPUT_FIRST_ROW}:$A${INPUT_LAST_ROW})'
    for coordinate in ("F2", "F3"):
        ws[coordinate].font = Font(name="Calibri", size=9, bold=True, color=GRAY)
        ws[coordinate].alignment = Alignment(horizontal="right", vertical="center")
    ws["G2"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    ws["G2"].border = Border(
        left=MEDIUM_BLUE,
        right=MEDIUM_BLUE,
        top=MEDIUM_BLUE,
        bottom=MEDIUM_BLUE,
    )
    for coordinate in ("G2", "G3"):
        ws[coordinate].font = Font(name="Calibri", size=9, bold=True, color=NAVY)
        ws[coordinate].alignment = Alignment(horizontal="center", vertical="center")

    mode_validation = DataValidation(type="list", formula1="=ModeLaporan")
    mode_validation.showDropDown = False
    mode_validation.promptTitle = "Mode laporan"
    mode_validation.prompt = "Pilih Harian atau Bulanan."
    mode_validation.errorTitle = "Pilihan tidak valid"
    mode_validation.error = "Pilih Harian atau Bulanan dari dropdown."
    mode_validation.showErrorMessage = True
    mode_validation.showInputMessage = True
    ws.add_data_validation(mode_validation)
    mode_validation.add("B2")

    day_validation = DataValidation(type="list", formula1="=DaftarHariBulan")
    day_validation.showDropDown = False
    day_validation.promptTitle = "Pilih tanggal"
    day_validation.prompt = "Daftar hanya menampilkan tanggal yang ada pada bulan terpilih."
    day_validation.errorTitle = "Tanggal tidak valid"
    day_validation.error = "Pilih tanggal dari dropdown."
    day_validation.showErrorMessage = True
    day_validation.showInputMessage = True
    ws.add_data_validation(day_validation)
    day_validation.add("E2")

    month_validation = DataValidation(type="list", formula1="=DaftarBulan")
    month_validation.showDropDown = False
    month_validation.promptTitle = "Pilih bulan"
    month_validation.prompt = "Pilih bulan laporan."
    month_validation.errorTitle = "Bulan tidak valid"
    month_validation.error = "Pilih bulan dari dropdown."
    month_validation.showErrorMessage = True
    month_validation.showInputMessage = True
    ws.add_data_validation(month_validation)
    month_validation.add("E3")

    year_validation = DataValidation(type="list", formula1="=DaftarTahun")
    year_validation.showDropDown = False
    year_validation.promptTitle = "Pilih tahun"
    year_validation.prompt = "Pilih tahun laporan (2020-2040)."
    year_validation.errorTitle = "Tahun tidak valid"
    year_validation.error = "Pilih tahun dari dropdown."
    year_validation.showErrorMessage = True
    year_validation.showInputMessage = True
    ws.add_data_validation(year_validation)
    year_validation.add("G2")

    headers_top = ["UPT / Kecamatan", "Jumlah Petugas", "Target WP", "Realisasi", None, "Persentase", "Keterangan"]
    headers_bottom = [None, None, None, "Jumlah SPPT", "Jumlah RP", None, None]
    for col, value in enumerate(headers_top, 1):
        ws.cell(6, col, value)
    for col, value in enumerate(headers_bottom, 1):
        ws.cell(7, col, value)

    for merge_range in ("A6:A7", "B6:B7", "C6:C7", "D6:E6", "F6:F7", "G6:G7"):
        ws.merge_cells(merge_range)
    for row in range(6, 8):
        for col in range(1, 8):
            header_style(ws.cell(row, col))
    ws.row_dimensions[6].height = 24
    ws.row_dimensions[7].height = 24

    report_rows = []
    current_row = 8
    for upt, districts in UPT_DATA:
        upt_row = current_row
        ws.cell(upt_row, 1, upt)
        district_rows = []
        for district in districts:
            current_row += 1
            district_row = current_row
            district_rows.append(district_row)
            report_rows.append(district_row)
            ws.cell(district_row, 1, district)
            ws.cell(district_row, 2, report_sum_formula("Jumlah Petugas", district_row))
            ws.cell(district_row, 3, report_sum_formula("Target WP", district_row))
            ws.cell(district_row, 4, report_sum_formula("Realisasi SPPT", district_row))
            ws.cell(district_row, 5, report_sum_formula("Realisasi RP", district_row))
            ws.cell(district_row, 6, f'=IFERROR(D{district_row}/C{district_row},0)')
            ws.cell(district_row, 7, report_note_formula(district_row))

        first_child, last_child = district_rows[0], district_rows[-1]
        for col in range(2, 6):
            letter = ws.cell(upt_row, col).column_letter
            ws.cell(upt_row, col, f'=SUM({letter}{first_child}:{letter}{last_child})')
        ws.cell(upt_row, 6, f'=IFERROR(D{upt_row}/C{upt_row},0)')
        ws.cell(upt_row, 7, f'="Total "&A{upt_row}')

        for col in range(1, 8):
            cell = ws.cell(upt_row, col)
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            cell.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
            cell.border = Border(
                left=THIN_GRAY, right=THIN_GRAY, top=MEDIUM_BLUE, bottom=THIN_GRAY
            )
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        current_row += 1

    grand_total_row = current_row
    ws.cell(grand_total_row, 1, "GRAND TOTAL")
    upt_rows = [8, 12, 16, 20, 24, 28, 32]
    for col in range(2, 6):
        letter = ws.cell(grand_total_row, col).column_letter
        refs = ",".join(f"{letter}{row}" for row in upt_rows)
        ws.cell(grand_total_row, col, f"=SUM({refs})")
    ws.cell(grand_total_row, 6, f'=IFERROR(D{grand_total_row}/C{grand_total_row},0)')
    ws.cell(grand_total_row, 7, '=IF($B$2="Harian","Total harian","Total bulanan")')

    for row in report_rows:
        for col in range(1, 8):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=WHITE)
            cell.font = Font(name="Calibri", size=10, color="222222")
            cell.border = Border(
                left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
            )
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 7))
        ws.cell(row, 1).alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 22

    for col in range(1, 8):
        cell = ws.cell(grand_total_row, col)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.border = Border(
            left=MEDIUM_BLUE,
            right=MEDIUM_BLUE,
            top=MEDIUM_BLUE,
            bottom=MEDIUM_BLUE,
        )
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[grand_total_row].height = 26

    for row in range(8, grand_total_row + 1):
        for col in (2, 3, 4):
            ws.cell(row, col).number_format = '#,##0'
        ws.cell(row, 5).number_format = '"Rp" #,##0'
        ws.cell(row, 6).number_format = "0.00%"
        for col in range(2, 7):
            ws.cell(row, col).alignment = Alignment(
                horizontal="right", vertical="center", wrap_text=False
            )

    ws.conditional_formatting.add(
        f"F8:F{grand_total_row}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="F8696B",
            mid_type="num",
            mid_value=0.75,
            mid_color="FFEB84",
            end_type="num",
            end_value=1,
            end_color="63BE7B",
        ),
    )
    ws.conditional_formatting.add(
        f"F8:F{grand_total_row}",
        CellIsRule(
            operator="greaterThan",
            formula=["1"],
            fill=PatternFill("solid", fgColor="A9D18E"),
        ),
    )

    widths = {"A": 25, "B": 24, "C": 15, "D": 17, "E": 21, "F": 14, "G": 36}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:7"
    ws.print_area = f"A1:G{grand_total_row}"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.oddFooter.left.text = "Laporan UPT"
    ws.oddFooter.center.text = "Halaman &P dari &N"
    ws.oddFooter.right.text = "Dicetak &D"


def setup_instructions_sheet(ws):
    ws.title = "PETUNJUK"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "PETUNJUK PENGGUNAAN"
    title_style(ws["A1"], 16)
    ws.row_dimensions[1].height = 30

    instructions = [
        ("1", "Masukkan data", "Buka sheet INPUT HARIAN dan isi sel berwarna kuning."),
        ("2", "Satu baris satu catatan", "Setiap baris mewakili satu kecamatan pada satu tanggal."),
        ("3", "Kolom otomatis", "Kolom Hari dan UPT terisi otomatis; tidak perlu diketik."),
        ("4", "Hindari data ganda", "Tanggal + Kecamatan yang sama akan ditandai merah. Jika tetap dibiarkan, nilainya dijumlahkan pada laporan."),
        ("5", "Pilih tanggal", "Gunakan dropdown Tanggal, Bulan, dan Tahun. Dropdown Tanggal otomatis berhenti pada hari terakhir bulan terpilih."),
        ("6", "Filter harian", "Pilih mode Harian. Saat Tanggal, Bulan, atau Tahun diganti, tabel menampilkan data tepat pada tanggal tersebut."),
        ("7", "Filter bulanan", "Pilih mode Bulanan. Laporan menjumlahkan seluruh data dalam Bulan dan Tahun yang dipilih."),
        ("8", "Keterangan bulanan", "Pada mode Bulanan, kolom Keterangan menampilkan keterangan terakhir yang diisi untuk tiap kecamatan."),
        ("9", "Edit laporan", "Seluruh sheet LAPORAN dapat diedit. Hindari menimpa sel rumus jika ingin perhitungan dan pergantian tanggal tetap otomatis."),
        ("10", "Tambah data", "Gunakan baris kosong berikutnya di INPUT HARIAN. Tersedia hingga 10.000 baris input."),
    ]

    for col, value in enumerate(["No.", "Bagian", "Cara Pakai"], 1):
        ws.cell(3, col, value)
        header_style(ws.cell(3, col))

    for row, (number, section, description) in enumerate(instructions, 4):
        ws.cell(row, 1, number)
        ws.cell(row, 2, section)
        ws.cell(row, 3, description)
        for col in range(1, 4):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
            )
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else PALE_BLUE)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 2).font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        ws.row_dimensions[row].height = 34

    ws["A15"] = "Catatan"
    ws["A15"].font = Font(name="Calibri", size=10, bold=True, color=NAVY)
    ws["B15"] = "Jumlah Petugas, Target WP, SPPT, dan Rupiah dijumlahkan pada mode Bulanan. Persentase = Realisasi SPPT / Target WP."
    ws.merge_cells("B15:F15")
    ws["B15"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B15"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    ws.row_dimensions[15].height = 36

    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 80
    for col in ("D", "E", "F"):
        ws.column_dimensions[col].width = 12


def build_workbook():
    if not SOURCE.exists():
        raise FileNotFoundError(f"File sumber tidak ditemukan: {SOURCE}")

    # Load the source once so a corrupt/non-XLSX input fails before output is built.
    source_wb = load_workbook(SOURCE, read_only=True, data_only=False)
    source_wb.close()

    wb = Workbook()
    report_ws = wb.active
    input_ws = wb.create_sheet()
    instructions_ws = wb.create_sheet()
    reference_ws = wb.create_sheet()

    setup_reference_sheet(reference_ws)
    add_defined_name(wb, "DaftarKecamatan", "'REFERENSI'!$A$2:$A$22")
    add_defined_name(wb, "ModeLaporan", "'REFERENSI'!$D$2:$D$3")
    add_defined_name(wb, "DaftarBulan", "'REFERENSI'!$I$2:$I$13")
    add_defined_name(wb, "DaftarTahun", "'REFERENSI'!$J$2:$J$22")
    add_defined_name(
        wb,
        "DaftarHariBulan",
        "OFFSET('REFERENSI'!$H$2,0,0,"
        "DAY(EOMONTH(DATE('LAPORAN'!$G$2,"
        "MATCH('LAPORAN'!$E$3,'REFERENSI'!$I$2:$I$13,0),1),0)),1)",
    )

    setup_report_sheet(report_ws)
    setup_input_sheet(input_ws)
    setup_instructions_sheet(instructions_ws)

    reference_ws.sheet_state = "hidden"
    wb.active = 0
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    wb.properties.title = "Laporan Harian dan Bulanan UPT"
    wb.properties.subject = "Input harian dengan filter tanggal dan bulan"
    wb.properties.creator = "OpenAI Codex"
    wb.properties.description = (
        "Workbook laporan UPT berbasis tabel input harian dengan rekap otomatis."
    )
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    result = build_workbook()
    print(result)
