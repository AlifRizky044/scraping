from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import buat_laporan_upt_harian as base


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / "LAPORAN UPT BIASA.xlsx"
FIRST_ROW = 5
LAST_ROW = 10004

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
PALE_RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "666666"
THIN = Side(style="thin", color="B7B7B7")
MEDIUM = Side(style="medium", color=BLUE)


def border(style=THIN):
    return Border(left=style, right=style, top=style, bottom=style)


def style_title(cell, size=16):
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Calibri", size=size, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_header(cell):
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border()


def headers():
    return [
        "Tanggal",
        "Hari",
        "UPT",
        "Kecamatan",
        "Jumlah Petugas",
        "Target WP",
        "Realisasi SPPT",
        "Realisasi RP",
        "Keterangan",
    ]


def setup_input(ws):
    ws.title = "KECAMATAN"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:I1")
    ws["A1"] = "INPUT HARIAN PER KECAMATAN"
    style_title(ws["A1"])
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:I2")
    ws["A2"] = (
        "Isi satu baris untuk satu kecamatan pada satu tanggal. "
        "Gunakan tombol filter pada kolom Tanggal untuk melihat tanggal tertentu."
    )
    ws["A2"].fill = PatternFill("solid", fgColor=PALE_GREEN)
    ws["A2"].font = Font(name="Calibri", size=10, bold=True, color="375623")
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28

    for col, value in enumerate(headers(), 1):
        ws.cell(4, col, value)
        style_header(ws.cell(4, col))
    ws.row_dimensions[4].height = 34

    helper_row = 2
    for upt, districts in base.UPT_DATA:
        for district in districts:
            ws.cell(helper_row, 10, district)
            ws.cell(helper_row, 11, upt)
            helper_row += 1
    ws.column_dimensions["J"].hidden = True
    ws.column_dimensions["K"].hidden = True

    note = Comment(
        "Pilih tanggal, pilih kecamatan, lalu isi angka pada kolom kuning.",
        "Laporan UPT",
    )
    for row in range(FIRST_ROW, LAST_ROW + 1):
        ws.cell(
            row,
            2,
            f'=IF(A{row}="","",CHOOSE(WEEKDAY(A{row},2),"Senin","Selasa","Rabu","Kamis","Jumat","Sabtu","Minggu"))',
        )
        ws.cell(
            row,
            3,
            f'=IFERROR(VLOOKUP(D{row},$J$2:$K$22,2,FALSE),"")',
        )
        for col in range(1, 10):
            cell = ws.cell(row, col)
            cell.border = border()
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 9))
            cell.fill = PatternFill(
                "solid", fgColor=PALE_BLUE if col in (2, 3) else PALE_YELLOW
            )
        ws.cell(row, 1).number_format = "dd mmm yyyy"
        for col in (5, 6, 7):
            ws.cell(row, col).number_format = "#,##0"
        ws.cell(row, 8).number_format = '"Rp" #,##0'
    ws["A5"].comment = note

    date_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2040,12,31)",
        allow_blank=True,
    )
    date_validation.showErrorMessage = True
    date_validation.errorTitle = "Tanggal tidak valid"
    date_validation.error = "Masukkan tanggal antara 2020 dan 2040."
    ws.add_data_validation(date_validation)
    date_validation.add(f"A{FIRST_ROW}:A{LAST_ROW}")

    district_validation = DataValidation(
        type="list", formula1="=$J$2:$J$22", allow_blank=True
    )
    district_validation.showDropDown = False
    district_validation.showErrorMessage = True
    district_validation.errorTitle = "Kecamatan tidak valid"
    district_validation.error = "Pilih kecamatan dari dropdown."
    ws.add_data_validation(district_validation)
    district_validation.add(f"D{FIRST_ROW}:D{LAST_ROW}")

    integer_validation = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True
    )
    ws.add_data_validation(integer_validation)
    integer_validation.add(f"E{FIRST_ROW}:G{LAST_ROW}")

    money_validation = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True
    )
    ws.add_data_validation(money_validation)
    money_validation.add(f"H{FIRST_ROW}:H{LAST_ROW}")

    duplicate_formula = (
        f'=AND($A{FIRST_ROW}<>"",$D{FIRST_ROW}<>"",'
        f'COUNTIFS($A${FIRST_ROW}:$A${LAST_ROW},$A{FIRST_ROW},'
        f'$D${FIRST_ROW}:$D${LAST_ROW},$D{FIRST_ROW})>1)'
    )
    ws.conditional_formatting.add(
        f"A{FIRST_ROW}:I{LAST_ROW}",
        FormulaRule(
            formula=[duplicate_formula],
            fill=PatternFill("solid", fgColor=PALE_RED),
            font=Font(color="C00000"),
        ),
    )

    widths = {
        "A": 15,
        "B": 13,
        "C": 12,
        "D": 23,
        "E": 16,
        "F": 14,
        "G": 17,
        "H": 20,
        "I": 38,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f"A4:I{LAST_ROW}"


def setup_source(ws):
    ws.title = "SUMBER DATA"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:I1")
    ws["A1"] = "SUMBER DATA OTOMATIS"
    style_title(ws["A1"])
    ws.merge_cells("A2:I2")
    ws["A2"] = (
        "Sheet ini terisi otomatis dari KECAMATAN. Jangan input manual di sini."
    )
    ws["A2"].fill = PatternFill("solid", fgColor="E7E6E6")
    ws["A2"].font = Font(italic=True, color=GRAY)

    for col, value in enumerate(headers(), 1):
        ws.cell(4, col, value)
        style_header(ws.cell(4, col))

    for row in range(FIRST_ROW, LAST_ROW + 1):
        for col in range(1, 10):
            source = ws.cell(row, col).column_letter
            cell = ws.cell(
                row,
                col,
                f'=IF(KECAMATAN!{source}{row}="","",KECAMATAN!{source}{row})',
            )
            cell.border = border()
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 9))
        ws.cell(row, 1).number_format = "dd mmm yyyy"
        for col in (5, 6, 7):
            ws.cell(row, col).number_format = "#,##0"
        ws.cell(row, 8).number_format = '"Rp" #,##0'

    widths = {
        "A": 15,
        "B": 13,
        "C": 12,
        "D": 23,
        "E": 16,
        "F": 14,
        "G": 17,
        "H": 20,
        "I": 38,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f"A4:I{LAST_ROW}"


def report_sum_formula(metric_col, row):
    value_range = f"'SUMBER DATA'!${metric_col}${FIRST_ROW}:${metric_col}${LAST_ROW}"
    district_range = f"'SUMBER DATA'!$D${FIRST_ROW}:$D${LAST_ROW}"
    date_range = f"'SUMBER DATA'!$A${FIRST_ROW}:$A${LAST_ROW}"
    return (
        f'=IF($B$3="","",IF($B$2="Harian",'
        f'IFERROR(SUMIFS({value_range},{district_range},$A{row},{date_range},$B$3),0),'
        f'IFERROR(SUMIFS({value_range},{district_range},$A{row},'
        f'{date_range},">="&EOMONTH($B$3,-1)+1,'
        f'{date_range},"<="&EOMONTH($B$3,0)),0)))'
    )


def report_note_formula(row):
    date_range = f"'SUMBER DATA'!$A${FIRST_ROW}:$A${LAST_ROW}"
    district_range = f"'SUMBER DATA'!$D${FIRST_ROW}:$D${LAST_ROW}"
    note_range = f"'SUMBER DATA'!$I${FIRST_ROW}:$I${LAST_ROW}"
    return (
        f'=IF($B$2="Harian",'
        f'IFERROR(LOOKUP(2,1/(({district_range}=$A{row})*({date_range}=$B$3)*'
        f'({note_range}<>"")),{note_range}),""),'
        f'IFERROR(LOOKUP(2,1/(({district_range}=$A{row})*'
        f'({date_range}>=EOMONTH($B$3,-1)+1)*'
        f'({date_range}<=EOMONTH($B$3,0))*({note_range}<>"")),{note_range}),""))'
    )


def setup_report(ws):
    ws.title = "LAPORAN"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    ws.merge_cells("A1:G1")
    ws["A1"] = "LAPORAN HARIAN / BULANAN UPT"
    style_title(ws["A1"], 17)
    ws.row_dimensions[1].height = 32

    ws["A2"] = "Mode Laporan"
    ws["A3"] = "Tanggal Acuan"
    ws["A4"] = "Periode Tampil"
    for row in range(2, 5):
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row, 1).font = Font(bold=True, color=NAVY)
        ws.cell(row, 1).border = border()

    ws["B2"] = "Harian"
    ws["B3"] = "=TODAY()"
    ws["B3"].number_format = "dd mmmm yyyy"
    month_text = 'CHOOSE(MONTH($B$3),"Januari","Februari","Maret","April","Mei","Juni","Juli","Agustus","September","Oktober","November","Desember")'
    day_text = 'CHOOSE(WEEKDAY($B$3),"Minggu","Senin","Selasa","Rabu","Kamis","Jumat","Sabtu")'
    ws["B4"] = (
        f'=IF($B$2="Harian",{day_text}&", "&TEXT($B$3,"dd")&" "&'
        f'{month_text}&" "&YEAR($B$3),{month_text}&" "&YEAR($B$3))'
    )
    ws.merge_cells("B4:G4")
    for coordinate in ("B2", "B3"):
        ws[coordinate].fill = PatternFill("solid", fgColor=PALE_YELLOW)
        ws[coordinate].font = Font(bold=True, color=NAVY)
        ws[coordinate].alignment = Alignment(horizontal="center", vertical="center")
        ws[coordinate].border = border(MEDIUM)
    ws["B4"].fill = PatternFill("solid", fgColor=PALE_GREEN)
    ws["B4"].font = Font(bold=True, color="375623")
    ws["B4"].alignment = Alignment(vertical="center")

    mode_validation = DataValidation(type="list", formula1='"Harian,Bulanan"')
    mode_validation.showDropDown = False
    ws.add_data_validation(mode_validation)
    mode_validation.add("B2")
    date_validation = DataValidation(
        type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2040,12,31)"
    )
    ws.add_data_validation(date_validation)
    date_validation.add("B3")

    top_headers = [
        "UPT / Kecamatan", "Jumlah Petugas", "Target WP", "Realisasi", None,
        "Persentase", "Keterangan",
    ]
    bottom_headers = [None, None, None, "Jumlah SPPT", "Jumlah RP", None, None]
    for col, value in enumerate(top_headers, 1):
        ws.cell(6, col, value)
    for col, value in enumerate(bottom_headers, 1):
        ws.cell(7, col, value)
    for merged in ("A6:A7", "B6:B7", "C6:C7", "D6:E6", "F6:F7", "G6:G7"):
        ws.merge_cells(merged)
    for row in range(6, 8):
        for col in range(1, 8):
            style_header(ws.cell(row, col))

    district_rows = []
    current_row = 8
    upt_rows = []
    for upt, districts in base.UPT_DATA:
        upt_row = current_row
        upt_rows.append(upt_row)
        ws.cell(upt_row, 1, upt)
        children = []
        for district in districts:
            current_row += 1
            row = current_row
            children.append(row)
            district_rows.append(row)
            ws.cell(row, 1, district)
            ws.cell(row, 2, report_sum_formula("E", row))
            ws.cell(row, 3, report_sum_formula("F", row))
            ws.cell(row, 4, report_sum_formula("G", row))
            ws.cell(row, 5, report_sum_formula("H", row))
            ws.cell(row, 6, f"=IFERROR(D{row}/C{row},0)")
            ws.cell(row, 7, report_note_formula(row))
        for col in range(2, 6):
            letter = ws.cell(upt_row, col).column_letter
            ws.cell(upt_row, col, f"=SUM({letter}{children[0]}:{letter}{children[-1]})")
        ws.cell(upt_row, 6, f"=IFERROR(D{upt_row}/C{upt_row},0)")
        ws.cell(upt_row, 7, f'="Total "&A{upt_row}')
        current_row += 1

    grand_row = current_row
    ws.cell(grand_row, 1, "GRAND TOTAL")
    for col in range(2, 6):
        letter = ws.cell(grand_row, col).column_letter
        ws.cell(grand_row, col, "=SUM(" + ",".join(f"{letter}{row}" for row in upt_rows) + ")")
    ws.cell(grand_row, 6, f"=IFERROR(D{grand_row}/C{grand_row},0)")
    ws.cell(grand_row, 7, '=IF(B2="Harian","Total harian","Total bulanan")')

    for row in range(8, grand_row + 1):
        is_total = row in upt_rows or row == grand_row
        for col in range(1, 8):
            cell = ws.cell(row, col)
            cell.border = border()
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 7))
            if is_total:
                color = NAVY if row == grand_row else LIGHT_BLUE
                font_color = WHITE if row == grand_row else NAVY
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=True, color=font_color)
        for col in (2, 3, 4):
            ws.cell(row, col).number_format = "#,##0"
        ws.cell(row, 5).number_format = '"Rp" #,##0'
        ws.cell(row, 6).number_format = "0.00%"

    ws.conditional_formatting.add(
        f"F8:F{grand_row}",
        ColorScaleRule(
            start_type="num", start_value=0, start_color="F8696B",
            mid_type="num", mid_value=0.75, mid_color="FFEB84",
            end_type="num", end_value=1, end_color="63BE7B",
        ),
    )
    widths = {"A": 25, "B": 16, "C": 15, "D": 17, "E": 21, "F": 14, "G": 36}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.print_area = f"A1:G{grand_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def setup_instructions(ws):
    ws.title = "PETUNJUK"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:C1")
    ws["A1"] = "CARA PAKAI"
    style_title(ws["A1"])
    rows = [
        ("1", "Input data hanya di sheet KECAMATAN."),
        ("2", "Satu baris adalah satu kecamatan pada satu tanggal."),
        ("3", "Klik panah filter di kolom Tanggal untuk menampilkan tanggal tertentu."),
        ("4", "Sheet SUMBER DATA menyalin input secara otomatis; jangan diketik manual."),
        ("5", "Sheet LAPORAN membuat rekap Harian atau Bulanan dari SUMBER DATA."),
        ("6", "Tanggal + Kecamatan yang sama ditandai merah karena akan dijumlahkan ganda."),
    ]
    for row, (number, text) in enumerate(rows, 3):
        ws.cell(row, 1, number)
        ws.cell(row, 2, text)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row, 1).border = border()
        ws.cell(row, 2).border = border()
        ws.row_dimensions[row].height = 30
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 85
    ws.column_dimensions["C"].width = 10


def build_workbook():
    wb = Workbook()
    input_ws = wb.active
    source_ws = wb.create_sheet()
    report_ws = wb.create_sheet()
    instructions_ws = wb.create_sheet()

    setup_input(input_ws)
    setup_source(source_ws)
    setup_report(report_ws)
    setup_instructions(instructions_ws)

    input_ws.sheet_properties.tabColor = "70AD47"
    source_ws.sheet_properties.tabColor = "A5A5A5"
    report_ws.sheet_properties.tabColor = "5B9BD5"
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_workbook())
